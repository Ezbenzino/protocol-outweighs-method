# -*- coding: utf-8 -*-
"""
v20 -> v21（2026-10-02）：改投 CMIG + 第三轮投稿前审计修复。

背景：原目标期刊 Computers in Biology and Medicine 已于 2025-11-17 被 Web of Science
（SCIE）剔除，改投 Computerized Medical Imaging and Graphics（CMIG）（2026-10-02）。

输入（只读，不改）：outputs/paper/论文_协议效应_v20_CIBM.docx
                    outputs/paper/论文_协议效应_v18_补充材料.docx
输出：            outputs/paper/论文_协议效应_v21_CMIG.docx
                    outputs/paper/论文_协议效应_v21_补充材料.docx
用法：python scripts/revise_paper_v21.py [输出目录]   # 不给参数就写到 outputs/paper/

每一处修改都先断言"旧文本在该段中恰好出现一次"；任何一条断言失败就整体中止、不写文件。
本脚本不改动任何实验数字；所有新出现的数字都来自已有 JSON（见各条注释）。
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'outputs' / 'paper'
OUTDIR = Path(sys.argv[1]) if len(sys.argv) > 1 else PAPER
SRC_MAIN = PAPER / '论文_协议效应_v20_CIBM.docx'
SRC_SUPP = PAPER / '论文_协议效应_v18_补充材料.docx'
DST_MAIN = OUTDIR / '论文_协议效应_v21_CMIG.docx'
DST_SUPP = OUTDIR / '论文_协议效应_v21_补充材料.docx'

LOG = []


# ----------------------------------------------------------------- helpers
def find_para(doc, prefix, contains=None):
    hits = [p for p in doc.paragraphs
            if p.text.startswith(prefix) and (contains is None or contains in p.text)]
    assert len(hits) == 1, f'find_para({prefix[:50]!r}) -> {len(hits)} hits'
    return hits[0]


def _check_runs(p, tag):
    joined = ''.join(r.text for r in p.runs)
    assert joined == p.text, f'[{tag}] paragraph has text outside plain runs; refuse to edit'


def replace(p, old, new, tag):
    """在单个 run 内做替换（不跨 run，避免破坏斜体等格式）。"""
    _check_runs(p, tag)
    full = p.text
    assert full.count(old) == 1, f'[{tag}] expected 1 x {old[:60]!r}, found {full.count(old)}'
    for r in p.runs:
        if old in r.text:
            assert r.text.count(old) == 1
            r.text = r.text.replace(old, new)
            break
    else:
        raise AssertionError(f'[{tag}] {old[:60]!r} spans several runs')
    assert p.text == full.replace(old, new), f'[{tag}] post-check failed'
    LOG.append(tag)


def set_single_run_text(p, new, tag):
    _check_runs(p, tag)
    assert len(p.runs) == 1, f'[{tag}] expected a single run, found {len(p.runs)}'
    p.runs[0].text = new
    LOG.append(tag)


def clone_after(anchor, template, segments, tag):
    """在 anchor 段之后插入新段：段落属性与首个 run 的字符格式取自 template。
    segments = [(text, {'italic':bool,'hl':bool}), ...]"""
    new_p = copy.deepcopy(template._p)
    proto = new_p.find(qn('w:r'))
    assert proto is not None, f'[{tag}] template has no run'
    proto = copy.deepcopy(proto)
    for child in list(new_p):
        if child.tag != qn('w:pPr'):
            new_p.remove(child)
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    for text, fmt in segments:
        r_el = copy.deepcopy(proto)
        for child in list(r_el):
            if child.tag != qn('w:rPr'):
                r_el.remove(child)
        new_p.append(r_el)
        run = para.runs[-1]
        run.text = text
        run.italic = True if fmt.get('italic') else None
        if fmt.get('hl'):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    LOG.append(tag)
    return para


def highlight_placeholder(p, placeholder, tag):
    """把段内某个占位符单独切成一个高亮 run（占位符须完整位于某一个 run 内）。"""
    _check_runs(p, tag)
    for r in p.runs:
        if placeholder in r.text:
            before, after = r.text.split(placeholder, 1)
            r.text = before
            # 依次在 r 之后插入 [placeholder][after]
            ph = copy.deepcopy(r._r); r._r.addnext(ph)
            tail = copy.deepcopy(r._r); ph.addnext(tail)
            from docx.text.run import Run
            ph_run, tail_run = Run(ph, p), Run(tail, p)
            ph_run.text = placeholder
            ph_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            tail_run.text = after
            LOG.append(tag)
            return
    raise AssertionError(f'[{tag}] placeholder not found in a single run')


I = {'italic': True}
N = {}
H = {'hl': True}

# =================================================================== MAIN
doc = docx.Document(str(SRC_MAIN))

# ---- A. 作者区：通讯作者邮箱占位符改为英文、高亮（作者本人填写）
p = find_para(doc, 'Yize Li')
replace(p, '*Correspondence: [PENDING: 通讯作者邮箱]',
        '*Corresponding author. E-mail address: [to be completed]', 'A1 corresponding-author line')
highlight_placeholder(p, '[to be completed]', 'A2 highlight e-mail placeholder')

# ---- B. 摘要：全部头条数字统一为保留测试集（论文声明的主推断）；
#         删除"protocol 主导 data-preparation"这一与 Table 5 矛盾的说法（偏移 26.36 本身属 data/training protocol）。
#      数字来源：analysis_test/symmetric_analysis_applied.json（28.58、9.99、3.12、1.87）、
#                analysis_scope/scope_analysis_merged2.json（22.51）、diag_offset_allarms.json（26.36）。
set_single_run_text(find_para(doc, 'Background and objective.'),
    'Background and objective. Reported Dice scores for lung nodule segmentation on LIDC-IDRI span 0.43 to 0.99, '
    'far more than the 0.01–0.03 gains individual methods claim. Every study chooses a label-fusion rule, a threshold, '
    'a patch-cropping rule and an evaluation extent; most do not state them. We ask how much of this spread comes from '
    'protocol rather than model, and whether probabilistic supervision helps.', 'B1 abstract background')
set_single_run_text(find_para(doc, 'Methods. On LIDC-IDRI'),
    'Methods. Four protocol factors — consensus level, binarisation threshold, evaluation extent and nodule-centred '
    'patch cropping — were controlled experimentally on 808 cross-validation cases and 202 held-out test cases. Four '
    'supervision targets, an SVLS comparator and five architectures were trained under an otherwise identical 2-D '
    'pipeline, and effects were compared with a repeated-training noise floor.', 'B2 abstract methods')
set_single_run_text(find_para(doc, 'Results. Widening the evaluation window'),
    'Results. On the held-out test set, changing only the evaluation target moved Dice by 28.58 points and only the '
    'threshold by 9.99; widening the evaluation window moved it by 22.51 points after both training-side confounds '
    'were removed, and an 8-pixel window offset cost a median of 26.36 points. The largest learning-design effect was '
    '3.12 points and the architecture range 1.87, against a 1.19-point noise floor; even the conservative '
    'field-of-view estimate exceeded these by 7.2× and 12.0×, respectively. Cross-validation reproduced the ordering. '
    'Soft supervision over rater votes conferred no measurable benefit. On seven external multi-rater tasks, the '
    'protocol effect fell as rater agreement rose.', 'B3 abstract results')
set_single_run_text(find_para(doc, 'Conclusions. On LIDC-IDRI'),
    'Conclusions. On LIDC-IDRI, the four protocol choices outweigh every learning-design choice measured, by three- '
    'to ninefold. Cross-study comparisons that omit them are not interpretable.', 'B4 abstract conclusions')

# ---- C. 引言
p = find_para(doc, 'A single well-cited example')
replace(p, 'compares its own result against five prior methods in one table whose rows are based on',
        'compares its own result against three prior methods in one table whose five rows are based on',
        'C1 NoduleNet: three prior methods (verified against arXiv:1907.11320 table)')
p = find_para(doc, '3. A direct comparison of protocol effect sizes')
replace(p, 'by 12.0× on the conservative estimate, with the same ordering reproduced on a held-out test set.',
        'by 12.0× on the conservative held-out-test estimate, with the same ordering reproduced in cross-validation.',
        'C2 contribution 3: ratios are test-set values')

# ---- C3. §2.4 指 Zhi 等人的图号写明归属，避免与本文 Fig. 9 混淆（2026-09-10 审计 §4.1 建议）
p = find_para(doc, 'Third, their protocol choices illustrate the reporting gap')
replace(p, '(their Section 3.6.2 and Fig. 9)', "(Zhi et al.'s Section 3.6.2 and Fig. 9)", 'C3a Zhi Fig. 9 attribution')
replace(p, '(their Section 3.6.2 and Fig. 11)', "(Zhi et al.'s Section 3.6.2 and Fig. 11)", 'C3b Zhi Fig. 11 attribution')

# ---- D. 相关工作：两处误引
p = find_para(doc, 'A parallel line of work preserves annotation ambiguity')
replace(p, 'in the model; Zhang et al. survey the field [17]. The Probabilistic U-Net [7] and its hierarchical extension, PHiSeg [18],',
        "in the model, or models each annotator's errors explicitly [17]. The Probabilistic U-Net [7], the hierarchical PHiSeg model [18],",
        'D1 [17] is a method paper, not a survey; PHiSeg is not an extension of the Probabilistic U-Net')

# ---- E. 方法
p = find_para(doc, 'On the held-out test set, all reported numbers are computed')
replace(p, 'On the held-out test set, all reported numbers are computed from the five fold models jointly:',
        'On the held-out test set, reported numbers are computed from the five fold models jointly unless stated otherwise:',
        'E1 test-set ensemble statement qualified')
replace(p, 'single-fold records are released alongside.',
        'single-fold records are released alongside. Two analyses use fewer folds: '
        "the ten-arm window-offset comparison of Section 4.10 uses each arm's fold-0 model, and the field-of-view arms "
        'of Sections 4.11 and 4.13 were evaluated on two folds (Section 5.5).',
        'E2 disclose fold-0 offset ladder and two-fold scope arms')
p = find_para(doc, 'The unit of analysis is the case, not the patch')
replace(p, 'The twenty pre-specified contrasts of Table 4 are corrected',
        'The twenty pre-specified paired contrasts of Section 4.3 — sixteen tabulated there and four reported in the text — are corrected',
        'E3 Holm family = 20 contrasts (Table 4 shows 16); Table 4 no longer first cited in 3.9')

# ---- F. 结果
anchor = find_para(doc, 'How much of this is arithmetic.')
# 数字来源：analysis_full/symmetric_analysis.json table1（tgtB: V2-V3 @0.5 = 11.42, @opt = 10.00；A–D 区间 9.86–13.81）
#           analysis_test/symmetric_analysis_applied.json table1（tgtB: 11.52 / 9.82）
clone_after(anchor, anchor, [
    ('Adjacent references, not only extreme ones. The effect does not depend on contrasting the union with the '
     'unanimous reference. Between ', N), ('V', I), (' ≥ 2 and ', N), ('V', I),
    (' ≥ 3 — adjacent levels, both in published use [2,3,6] — the majority-trained arm loses 11.42 points at a fixed '
     'threshold of 0.5 and 10.00 points at its calibrated thresholds in cross-validation (11.52 and 9.82 on the '
     'held-out test set), and every supervision arm in Table 3 loses between 9.86 and 13.81 points under one '
     'convention or the other. Two studies that report against ', N), ('V', I), (' ≥ 2 and ', N), ('V', I),
    (' ≥ 3 respectively are therefore separated by roughly ten Dice points before any difference in method is '
     'considered — about four times the largest supervision-target effect and ten times the architecture range.', N),
], 'F1 new paragraph: adjacent-reference effect (V>=2 vs V>=3)')

p = find_para(doc, 'Against the two references in common use, the effect is a well-powered null.')
replace(p, 'Case-level paired testing gives +0.06', 'Case-level paired testing (Table 4) gives +0.06',
        'F2a first in-text citation of Table 4 now follows Table 3')
p = find_para(doc, 'Consequently, the apparent advantage of')
replace(p, 'Decomposing the conventional A-versus-C comparison:',
        'Decomposing the conventional A-versus-C comparison into cells of the calibrated matrix (Fig. 4):',
        'F2b first in-text citation of Fig. 4 now precedes Fig. 5')
p = find_para(doc, 'Table 4.', contains='Case-level paired comparisons')
replace(p, 'Positive values favour the first arm.',
        'Positive values favour the first arm. p (Holm) is adjusted over all twenty pre-specified contrasts, including '
        'the four D − C contrasts reported in Section 4.3.', 'F2 Table 4 caption: Holm family')

p = find_para(doc, '4.5 Evaluation-protocol effects dominate')
replace(p, '4.5 Evaluation-protocol effects dominate data-preparation and learning-design effects',
        '4.5 Protocol effects dominate learning-design effects', 'F3 heading 4.5')

t5 = doc.tables[4]
r3, r4 = t5.rows[3], t5.rows[4]
assert r3.cells[0].text.startswith('Evaluation field of view, 128 → 512 px, both'), r3.cells[0].text
assert r4.cells[0].text.startswith('Evaluation target, V ≥ 1 to V ≥ 4'), r4.cells[0].text
r3._tr.addprevious(r4._tr)
LOG.append('F4 Table 5 rows put in ranked order (25.58 before 22.51)')

p = find_para(doc, 'We use the conservative figure for every comparative claim')
replace(p, 'the largest evaluation-protocol effect exceeds the largest learning-design effect',
        'the conservative field-of-view effect exceeds the largest learning-design effect', 'F5 ratio numerator named')
replace(p, 'every evaluation-protocol factor in Table 5 exceeds every data-preparation and learning-design factor.',
        'every evaluation-protocol factor in Table 5 exceeds every learning-design factor, as does the position prior '
        'installed by nodule-centred cropping.', 'F6 remove false claim (offset 26.36 is data/training protocol)')

p = find_para(doc, 'Table 8.', contains='Held-out test set')
replace(p, 'at the imported threshold.',
        'at the imported threshold. Ranges quoted in the text are computed from unrounded case-level means and can '
        'differ by 0.01 from differences between the rounded entries.', 'F7 Table 8 caption: rounding note (1.87 vs 1.88)')

p = find_para(doc, '4.10 Nodule-centred patch sampling')
replace(p, 'costs 27 Dice points at an 8-pixel offset', 'costs a median of 26 Dice points at an 8-pixel offset',
        'F8 heading 4.10 (median 26.36, not 27)')

# 4.10：B 臂曲线统一为五折集成（diag_offset.json：d0 0.8731 / d8 0.5730 / d56 0.0345），
#       与同句 d56、d40 配对检验同口径；原 0.8703/0.5953 是 fold 0 数值。
p = find_para(doc, 'For arm B the Dice against')
replace(p, 'For arm B the Dice against ', 'For arm B, averaged over its five fold models, the Dice against ',
        'F9 arm B offset curve: five-fold')
replace(p, '₂ falls from 0.8703 at ', '₂ falls from 0.8731 at ', 'F10 d0 five-fold 0.8731')
replace(p, ' = 0 to 0.5953 at ', ' = 0 to 0.5730 at ', 'F11 d8 five-fold 0.5730')
# 十臂比较：fold 0 模型（diag_offset_allarms.json），d56 区间含 SoftSeg 应为 66.9–86.4
p = find_para(doc, 'The effect is not specific to one architecture or one supervision target.')
replace(p, 'Ten arms trained without translation augmentation lose',
        "Evaluated with each arm's fold-0 model, the ten arms trained without translation augmentation lose",
        'F12 ten-arm ladder uses fold-0 models')
replace(p, 'between 82.3 and 86.4 points at ', 'between 66.9 and 86.4 points at ',
        'F13 d56 range includes SoftSeg (66.9)')
replace(p, 'consensus, soft and SVLS supervision.', 'consensus, soft, SVLS and SoftSeg supervision.',
        'F14 list SoftSeg among the ten arms')
p = find_para(doc, 'Fig. 7', contains='displaced from the lesion centroid')
replace(p, 'the augmented arm is flat across exactly that range.',
        "the augmented arm is flat across exactly that range. All curves use each arm's fold-0 model on the 76 "
        'eligible test-set nodules; the paired comparison of Section 4.10 averages each nodule over the five fold models.',
        'F15 Fig. 7 caption: fold-0')

p = find_para(doc, 'SoftSeg [8] combines an unbinarised soft target, a normalised ReLU final')
replace(p, 'reports a threshold sweep in 0.05 steps over [0, 1]',
        'reports a threshold sweep from 0.05 to 0.95 in steps of 0.05', 'F16 SoftSeg sweep range (verified in [8])')

p = find_para(doc, 'A single continuous quantity reconciles the datasets.')
replace(p, 'Spearman rho = -0.79', 'Spearman ρ = −0.79', 'F17a')
replace(p, 'yields rho = -0.86', 'yields ρ = −0.86', 'F17b')
replace(p, '(rho remains between -0.79 and -0.89)', '(ρ remains between −0.79 and −0.89)', 'F17c')
replace(p, '(0.5-0.8 points)', '(0.5–0.8 points)', 'F17d')
replace(p, '(prostate t1, 4.58x)', '(prostate t1, 4.58×)', 'F17e')
p = find_para(doc, 'Fig. 9', contains='Absolute protocol effect')
replace(p, '(Spearman rho = -0.86,', '(Spearman ρ = −0.86,', 'F18 Fig. 9 caption typography')

# ---- G. 讨论
p = find_para(doc, 'The central observation is arithmetic.')
replace(p, 'if the detector supplying its crop is off by two lesion radii, at 0.87 or 0.57.',
        'if the detector supplying its crop is off by 8 pixels — about 1.6 lesion radii — at 0.87 or 0.57.',
        'G1 offset = 8 px ≈ 1.6 radii (five-fold 0.8731→0.5730)')
replace(p, 'and the encoder learning rate (Section 4.8)', 'and the learning rate (Section 4.8)',
        'G2 the swept learning rate is the plain U-Net one')
p = find_para(doc, 'Five of the published differences cited in this paper')
set_single_run_text(p,
    'Several published differences cited in this paper are small against the protocol effects measured here. '
    'NoduleNet reports an architectural gain of 0.95 points over three prior methods [5]; the five best models of '
    'Zhi et al. span 0.31 points across a 558-fold capacity range [31]; and the supervision targets of this study '
    'span at most 2.65 points. Each is smaller, by a factor of at least 3.5, than every protocol effect measured '
    'here: 25.58 points for the consensus level, 9.42 for the threshold and 22.51 for the field of view. The overall '
    'spread of the reported LIDC literature, roughly 0.43 to 0.99, is by contrast of the same order as those protocol '
    'effects, which is consistent with unmatched protocols accounting for much of it.', 'G3 rewrite muddled paragraph')

p = find_para(doc, 'Two results in the wider literature support this reading.')
_check_runs(p, 'G4')
runs = p.runs
assert len(runs) == 9 and runs[1].text == 'p' and runs[3].text == 'V' and runs[5].text == 'G' and runs[7].text == 'G'
old0 = ('Two results in the wider literature support this reading. Müller et al. showed that label smoothing improves '
        'calibration and confidence estimates without improving accuracy, and that it erases information in exactly '
        'the regions where the target is already unambiguous [38]; the vote fraction ')
assert runs[0].text == old0, runs[0].text
runs[0].text = ('Two results in the wider literature bear on this reading. Müller et al. showed that label smoothing '
                'improves calibration by preventing over-confident outputs, at the price of erasing information from '
                'the logits about the similarity between classes [38]; the vote fraction ')
assert runs[4].text == '/4 is a data-driven label smoothing, and ', runs[4].text
runs[4].text = ('/4 is a data-driven label smoothing, and the band of partial agreement in which it suppresses '
                'confident outputs is where ')
assert runs[8].text.startswith('₄ are precisely the unambiguous regions. '), runs[8].text
runs[8].text = runs[8].text.replace('₄ are precisely the unambiguous regions. ', '₄ require a confident negative. ', 1)
LOG.append('G4 Müller [38] characterised as in its abstract (no "without improving accuracy")')

p = find_para(doc, 'The probabilistic multi-rater models most likely to be raised')
replace(p, 'the Probabilistic U-Net [7] and its hierarchical extension PHiSeg [18]',
        'the Probabilistic U-Net [7] and the hierarchical PHiSeg model [18]', 'G5 PHiSeg wording')

anchor = find_para(doc, 'A second recommendation follows from Section 4.10')
clone_after(anchor, anchor, [
    ('A third recommendation concerns measurement rather than evaluation. The fused reference is not only a scoring '
     'convention: the union reference covers 2.9 times the area of the unanimous one (Section 4.1), which for a '
     'lesion of mean size corresponds to roughly a 1.7-fold difference in equivalent diameter, and management '
     'recommendations for incidental pulmonary nodules change category at diameters of 6 and 8 mm [42]. Any '
     'segmentation-derived size or volume therefore inherits the fusion rule against which the segmenter was trained '
     'and validated, and that rule should be reported wherever such measurements are intended for clinical use.', N),
], 'G6 new paragraph: clinical consequence of the fusion rule')

# ---- H. 结论
p = find_para(doc, 'On LIDC-IDRI lung nodule segmentation, four evaluation-protocol choices')
set_single_run_text(p,
    'On LIDC-IDRI lung nodule segmentation, four protocol choices — the consensus level of the reference, the '
    'binarisation threshold, the evaluation extent and the patch-cropping rule — each move reported Dice by more '
    'than any learning-design choice measured on the same data.', 'H1 conclusion opening (category-consistent)')
p = find_para(doc, 'Widening the evaluation window from 128 to 512 pixels')
replace(p, 'Changing only the consensus level used as the reference costs 25.58 points, and the binarisation threshold alone 9.42.',
        'Changing only the consensus level used as the reference costs 28.58 points on the held-out test set (25.58 in '
        'cross-validation), and the binarisation threshold alone 9.99 (9.42).', 'H2 conclusion: label test vs CV')
p = find_para(doc, 'Finally, the dominance of protocol over method')
replace(p, '(Spearman rho = −0.86,', '(Spearman ρ = −0.86,', 'H3 typography')

# ---- I. 声明
p = find_para(doc, '[PENDING: 例如')
set_single_run_text(p, 'The authors declare that they have no known competing financial interests or personal '
                       'relationships that could have appeared to influence the work reported in this paper.',
                    'I1 competing interests (standard Elsevier wording; authors to confirm)')
p = find_para(doc, '[PENDING: 按 CRediT')
set_single_run_text(p, 'Yize Li: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, '
                       'Data curation, Visualization, Writing – original draft. Guanqun Sun: Supervision, Writing – '
                       'review & editing.', 'I2 CRediT (draft; authors to confirm)')
p = find_para(doc, '[PENDING: 有资助填资助号')
set_single_run_text(p, '[Funding statement to be completed: grant name and number, or "This research did not '
                       'receive any specific grant from funding agencies in the public, commercial, or not-for-profit '
                       'sectors."]', 'I3 funding placeholder (English)')
p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
p = find_para(doc, '[PENDING: 代码仓库 URL]')
replace(p, '[PENDING: 代码仓库 URL] The release', '[Code repository URL and Zenodo DOI to be inserted.] The release',
        'I4 data availability placeholder (English)')
replace(p, 'LIDC-IDRI is available from The Cancer Imaging Archive;',
        'LIDC-IDRI is available from The Cancer Imaging Archive [43];', 'I5 cite LIDC-IDRI dataset DOI (TCIA policy)')
highlight_placeholder(p, '[Code repository URL and Zenodo DOI to be inserted.]', 'I6 highlight repo placeholder')
p_ack = find_para(doc, '[PENDING]')
set_single_run_text(p_ack, 'The authors acknowledge the National Cancer Institute and the Foundation for the National '
                           'Institutes of Health, and their critical role in the creation of the free publicly '
                           'available LIDC/IDRI Database used in this study. We thank the organisers of the QUBIQ '
                           'challenge for making the multi-rater annotations available.',
                    'I7 acknowledgements (TCIA-required LIDC statement)')
title_tpl = find_para(doc, 'Acknowledgements')
ai_title = clone_after(p_ack, title_tpl, [
    ('Declaration of generative AI and AI-assisted technologies in the manuscript preparation process', N)],
    'I8 AI-declaration heading (Elsevier policy)')
clone_after(ai_title, p_ack, [
    ('During the preparation of this work the authors used Claude (Anthropic), Doubao (ByteDance) and Codex (OpenAI) '
     'in order to assist with writing and debugging analysis code, drafting and language editing of the manuscript, '
     'and cross-checking reported numbers against the analysis outputs. After using these tools, the authors reviewed '
     'and edited the content as needed and take full responsibility for the content of the published article.', N)],
    'I9 AI-declaration text (authors to confirm the tool list)')

# ---- J. 参考文献
p = find_para(doc, '[38] R. Müller')
replace(p, 'pp. 4696–4705', 'pp. 4694–4703', 'J1 [38] page range')
p41 = find_para(doc, '[41] A. Mehrtash')
p42 = clone_after(p41, p41, [(
    '[42] H. MacMahon, D.P. Naidich, J.M. Goo, K.S. Lee, A.N.C. Leung, J.R. Mayo, et al., Guidelines for management of '
    'incidental pulmonary nodules detected on CT images: from the Fleischner Society 2017, Radiology 284 (2017) '
    '228–243. https://doi.org/10.1148/radiol.2017161659.', N)], 'J2 add [42] Fleischner 2017 (Crossref-verified)')
clone_after(p42, p41, [(
    '[43] [dataset] S.G. Armato III, G. McLennan, L. Bidaut, M.F. McNitt-Gray, C.R. Meyer, A.P. Reeves, et al., Data '
    'from LIDC-IDRI, The Cancer Imaging Archive, 2015. https://doi.org/10.7937/K9/TCIA.2015.LO9QL9SX.', N)],
    'J3 add [43] LIDC-IDRI data citation (TCIA)')

DST_MAIN.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(DST_MAIN))

# =================================================================== SUPPLEMENT
sdoc = docx.Document(str(SRC_SUPP))
p = find_para(sdoc, 'First, the spread within this study')
replace(p, 'Both studies, independently, find the model axis to be worth less than half a point.',
        'Both studies, independently, find the differences among their strongest models to be worth less than half a point.',
        'S1 supplement: 0.96 is not < 0.5; the claim holds for the strongest models (0.19 / 0.31)')
sdoc.save(str(DST_SUPP))

print('applied', len(LOG), 'edits:')
for t in LOG:
    print('  -', t)
print('wrote', DST_MAIN)
print('wrote', DST_SUPP)
