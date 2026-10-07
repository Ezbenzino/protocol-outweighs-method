# -*- coding: utf-8 -*-
"""v22 投稿前自检（只读）。用法：python scripts/verify_paper_v22.py [docx路径]
检查：摘要词数≤250、Highlights≤85字符、正文引用按首次出现顺序编号、参考文献连续、
图表编号连续、无中文残留、列出所有待作者填写的占位符、Table 5 排序。"""
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'outputs' / 'paper' / '论文_协议效应_v22_CMIG.docx'
d = docx.Document(str(PATH))
P = [p.text for p in d.paragraphs]
ok = True


def fail(msg):
    global ok
    ok = False
    print('  ✗', msg)


# 1 摘要
i0 = P.index('Abstract')
abstract = [t for t in P[i0 + 1:i0 + 5]]
assert abstract[0].startswith('Background') and abstract[3].startswith('Conclusions')
n = sum(len(t.split()) for t in abstract)
print(f'[1] abstract words (whitespace tokens) = {n}')
if n > 250:
    fail('abstract exceeds 250 words')

# 2 Highlights
ih = P.index('Highlights')
hl = [t.lstrip('• ').strip() for t in P[ih + 1:ih + 6]]
print('[2] highlight lengths =', [len(h) for h in hl])
for h in hl:
    if len(h) > 85:
        fail(f'highlight > 85 chars: {h}')

# 3 引用顺序
iref = P.index('References')
seen, mx = [], 0
for t in P[:iref]:
    for m in re.finditer(r'\[(\d+(?:\s*[,–-]\s*\d+)*)\]', t):
        nums = []
        for part in re.split(r'\s*,\s*', m.group(1)):
            if re.search(r'[–-]', part):
                a, b = [int(x) for x in re.split(r'[–-]', part)]
                nums += list(range(a, b + 1))
            else:
                nums.append(int(part))
        for k in nums:
            if k not in seen:
                if k != mx + 1:
                    fail(f'first citation of [{k}] out of order (previous max {mx}): …{t[max(0, m.start() - 50):m.end()]}')
                seen.append(k)
                mx = max(mx, k)
refs = [t for t in P[iref + 1:] if re.match(r'^\[\d+\]', t)]
nums = [int(re.match(r'^\[(\d+)\]', t).group(1)) for t in refs]
print(f'[3] cited refs = {len(seen)} (max {mx}); reference list = {len(refs)} entries')
if nums != list(range(1, len(refs) + 1)):
    fail('reference list numbering not contiguous')
if sorted(seen) != nums:
    fail(f'cited set != listed set: missing {sorted(set(nums) - set(seen))}, extra {sorted(set(seen) - set(nums))}')

# 4 图表编号
figs = [int(m.group(1)) for t in P for m in [re.match(r'^Fig\. (\d+) —', t)] if m]
tabs = [int(m.group(1)) for t in P for m in [re.match(r'^Table (\d+)\.', t)] if m]
print('[4] figure captions', figs, '| table captions', tabs, '| tables in body', len(d.tables))
if figs != list(range(1, 10)) and sorted(figs) != list(range(1, 10)):
    fail('figure captions not 1–9')
if sorted(tabs) != list(range(1, 10)):
    fail('table captions not 1–9')

# 4b 图表按首次引用顺序编号（排除图注/表题本身与 "their Table/Fig"）
firsts = {'Fig': [], 'Table': []}
for t in P[:iref]:
    if re.match(r'^(Fig\. \d+ —|Table \d+\.)', t):
        continue
    for m in re.finditer(r'(Fig|Table)s?\.? (\d+)', t):
        ctx = t[max(0, m.start() - 45):m.start()]
        if re.search(r"[Tt]heir|Zhi et al\.’?'?s", ctx):
            continue  # 指 Zhi et al. 的图表，不是本文的
        kind, k = m.group(1), int(m.group(2))
        if k not in firsts[kind]:
            firsts[kind].append(k)
print('[4b] first-citation order  Fig:', firsts['Fig'], '| Table:', firsts['Table'])
for kind in firsts:
    if firsts[kind] != sorted(firsts[kind]):
        fail(f'{kind} numbers are not in order of first citation')

# 5 中文残留 / 占位符
cjk = [(i, t[:60]) for i, t in enumerate(P) if re.search(r'[一-鿿]', t)]
for c in cjk:
    fail(f'CJK text remains in paragraph {c[0]}: {c[1]}')
ph = [(i, m.group(0)) for i, t in enumerate(P) for m in re.finditer(r'\[(?:to be completed|PENDING[^\]]*|Funding statement to be completed[^\]]*|Code repository URL[^\]]*)\]', t)]
print('[5] placeholders left for the authors:')
for i, s in ph:
    print('     P%d: %s' % (i, s[:90]))

# 6 Table 5 排序
t5 = d.tables[4]
vals = []
for r in t5.rows[1:]:
    m = re.match(r'([\d.]+)', r.cells[2].text.strip())
    vals.append((float(m.group(1)), r.cells[0].text[:45], r.cells[1].text))
effects = [v for v in vals if v[2] != 'Noise floor']
print('[6] Table 5 order:', [v[0] for v in vals])
if [v[0] for v in effects] != sorted([v[0] for v in effects], reverse=True):
    fail('Table 5 effects not in descending order')

# 7 已知易错表述
bad = ['Spearman rho', 'rho =', 'five prior methods', 'survey the field', 'hierarchical extension',
       'two lesion radii', 'encoder learning rate', 'costs 27 Dice', 'data-preparation and learning-design effects',
       'exceeds every data-preparation', 'over [0, 1]', '4696–4705', 'between 82.3 and 86.4',
       # v22：已被重算替换的旧数字与已纠正的表述（2026-10-02 阈值统一与测试集 plain U-Net 修复）
       '22.51', '43.84', '26.36', '1.87', '12.0×', '0.7590', '0.3206', '0.5349', '0.2566', '0.0511', '14.84',
       '60.72', '55.40', '55.31', '0.8731', '0.5730', '0.0345', '66.9', '86.4', '260.3', '111.8', '5.8 mm',
       'below 0.1', 'near-diagonal', '183.2–188.7', 'twenty times', 'two upper intervals', '0.95 to 0.998',
       '10.4×', '23.4×', 'order of magnitude smaller than the evaluation-protocol', 'largest single factor',
       '7.2×', '7.2 times', 'reproduced the ordering', 'ordering is identical under both']
# 表格单元格同样检查旧数字
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            for b in ['22.51', '43.84', '26.36', '0.8121', '0.8532', '6.9–35.5']:
                if b in c.text:
                    fail(f'stale number "{b}" in a table cell: {c.text[:60]}')
# AI 声明须列出全部工具
ai = [t for t in P if t.startswith('During the preparation of this work the authors used')]
if not ai or 'Zhipu' not in ai[0]:
    fail('AI declaration does not list Zhipu (ChatGLM)')
for b in bad:
    for i, t in enumerate(P):
        if b in t:
            fail(f'stale wording "{b}" in P{i}')
print('\nRESULT:', 'ALL CHECKS PASSED' if ok else 'SOME CHECKS FAILED')
