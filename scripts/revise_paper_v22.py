# -*- coding: utf-8 -*-
"""
v21 -> v22（2026-10-02/03）：数据口径统一 + 全部插图重画。

本轮修改的来源（均为本仓库内的真实重算，见 SOURCE_OF_TRUTH.md 2026-10-02 条目）：
  1. 测试集 plain U-Net 换成选定学习率的 archPU1e3（run_fix_testset_plainunet_20261002.ps1）：
     架构极差 1.87 -> 1.19，比值 12.0x -> 18.7x，Table 8 的 Plain U-Net 行。
  2. 位置偏移与视野阶梯改用病例级验证集校准阈值（run_fix_offset_20261002.ps1；
     scope_analysis_merged2.json 以 analysis_full/symmetric_analysis.json 阈值重算）：
     26.36 -> 26.40；22.51 -> 22.29；43.84 -> 43.54；§4.10/4.11/4.13 全部相关数字。
  3. 正文与数据不符的表述：§3.1 测试集面积（旧版来自已废弃的去重口径）、§4.2 "case-level"
     面积（实为按 run 重复加权的行均值）、Fig. 1 直径（名义 0.6 mm/px 下为 4.9 mm）、
     Fig. 2/4/5 图注、§4.11 每窗最优阈值区间、两处"一个数量级"与"最大单一因素"的过度表述。
  4. 新增披露：名义像素间距 0.6 mm 与 20 像素纳入阈值（§3.1、Table 1）。
  5. AI 工具声明补充 ChatGLM（智谱 AI）。
  6. 9 张插图全部替换为 scripts/make_figures_v22.py 的输出（outputs/figures/v22/）。
  7. 补充材料：删除 9 个未被正文引用、却仍打包在 docx 里的旧插图（孤立图片关系）。

输入（只读）：outputs/paper/论文_协议效应_v21_CMIG.docx、论文_协议效应_v21_补充材料.docx、
              Highlights_v21.docx、Cover_letter_CMIG_v21.docx、outputs/figures/v22/Fig*.png
输出：        outputs/paper/论文_协议效应_v22_CMIG.docx、论文_协议效应_v22_补充材料.docx、
              Highlights_v22.docx、Cover_letter_CMIG_v22.docx
用法：python scripts/revise_paper_v22.py [输出目录]

每一处修改都先断言"旧文本在该段中恰好出现一次"；任何断言失败就整体中止、不写文件。
新数字全部从 JSON 现取并与下方文本中的字面量交叉断言，防止手抄错误。
"""
import json
import re
import struct
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'outputs' / 'paper'
FIGS = ROOT / 'outputs' / 'figures' / 'v22'
OUTDIR = Path(sys.argv[1]) if len(sys.argv) > 1 else PAPER
SRC = {k: PAPER / v for k, v in {
    'main': '论文_协议效应_v21_CMIG.docx', 'supp': '论文_协议效应_v21_补充材料.docx',
    'hl': 'Highlights_v21.docx', 'cover': 'Cover_letter_CMIG_v21.docx'}.items()}
DST = {k: OUTDIR / v for k, v in {
    'main': '论文_协议效应_v22_CMIG.docx', 'supp': '论文_协议效应_v22_补充材料.docx',
    'hl': 'Highlights_v22.docx', 'cover': 'Cover_letter_CMIG_v22.docx'}.items()}
LOG = []


def J(*p):
    with open(ROOT.joinpath(*p), encoding='utf-8') as f:
        return json.load(f)


# ------------------------------------------------------------- 数字：从 JSON 现取，并与文本字面量互相核对
T = J('outputs', 'analysis_test', 'symmetric_analysis_applied.json')['table1']
CV = J('outputs', 'analysis_full', 'symmetric_analysis.json')
EU = J('outputs', 'analysis', 'effect_uncertainty.json')
OFF = J('outputs', 'analysis_scope', 'offset_paired_stats.json')
SC = J('outputs', 'analysis_scope', 'scope_analysis_merged2.json')
NF = J('outputs', 'analysis_full', 'noise_floor.json')


def seg(arm, name):
    r = [x for x in SC['table4_paired'] if x['arm'] == arm and x['level'] == 2 and x['segment'] == name]
    assert len(r) == 1
    return r[0]


def lad(arm, scope):
    return SC['table1_ladder'][f'{arm}_v2'][scope]


def f2(x):
    return f'{x:.2f}'


def f4(x):
    return f'{x:.4f}'


def sci(p):
    m, e = f'{p:.1e}'.split('e')
    sup = str.maketrans('-0123456789', '⁻⁰¹²³⁴⁵⁶⁷⁸⁹')
    return f'{m} × 10{str(int(e)).translate(sup)}'


HE = EU['headline_point_estimates']
OWN = EU['own_sample_bootstrap']
JR = EU['joint_ratio_bootstrap']
N = {
    'fov': f2(HE['fov_effect_pts']), 'fov_pi': f2(HE['fov_position_invariant_pts']),
    'arch': f2(HE['architecture_range_pts']), 'sup': f2(HE['supervision_range_pts']),
    'r_sup': f'{HE["ratio_protocol_supervision"]:.1f}', 'r_arch': f'{HE["ratio_protocol_architecture"]:.1f}',
    'r_pi_sup': f'{HE["ratio_position_invariant_supervision"]:.1f}',
    'r_pi_arch': f'{HE["ratio_position_invariant_architecture"]:.1f}',
    'off_med': f2(OFF['ten_arms_fold0']['median_drop8_pts']),
    'off_lo': f'{OFF["ten_arms_fold0"]["range_drop8_pts"][0]:.1f}', 'off_hi': f'{OFF["ten_arms_fold0"]["range_drop8_pts"][1]:.1f}',
    'off56_lo': f'{OFF["ten_arms_fold0"]["range_drop56_pts"][0]:.1f}', 'off56_hi': f'{OFF["ten_arms_fold0"]["range_drop56_pts"][1]:.1f}',
}
# 这些字面量出现在下面的新文本里；先与 JSON 对齐，任何一项不符立即中止
EXPECT = {'fov': '22.29', 'fov_pi': '43.54', 'arch': '1.19', 'sup': '3.12', 'r_sup': '7.1', 'r_arch': '18.7',
          'r_pi_sup': '13.9', 'r_pi_arch': '36.6', 'off_med': '26.40', 'off_lo': '7.1', 'off_hi': '35.8',
          'off56_lo': '67.8', 'off56_hi': '86.5'}
for k, v in EXPECT.items():
    assert N[k] == v, (k, N[k], v)
assert [round(x, 1) for x in JR['ratio_protocol_supervision']['ci95']] == [5.8, 8.7] or \
    JR['ratio_protocol_supervision']['ci95'] == [5.84, 8.65], JR
assert f'{JR["ratio_protocol_supervision"]["point"]:.1f}' == '7.0'
assert f'{JR["ratio_protocol_architecture"]["point"]:.1f}' == '15.6'
assert [f'{x:.1f}' for x in JR['ratio_protocol_architecture']['ci95']] == ['9.3', '30.9']
assert OWN['supervision_range_pts']['ci95'] == [2.53, 3.71] and OWN['architecture_range_pts']['ci95'] == [0.57, 1.96]
assert NF['fixed_0_5']['two_sigma_pts'] == 1.19
pf = OFF['paired_d40']
assert (f'{pf["mean_pts"]:+.2f}', f'{pf["ci95_pts"][0]:+.2f}', f'{pf["ci95_pts"][1]:+.2f}', f'{pf["dz"]:.2f}') == \
    ('+60.79', '+55.47', '+65.80', '2.58'), pf
cb = OFF['fivefold_curves']['tgtB']
assert (f4(cb['0']), f4(cb['8']), f4(cb['56'])) == ('0.8733', '0.5718', '0.0343'), cb
assert f'{8 / OFF["eligible_nodules"]["median_equivalent_radius_px"]:.1f}' == '1.6'
fs, fm, fb, fbg = seg('tgtBshift', '纯视野效应'), seg('tgtBshift_bg', '纯视野效应'), seg('tgtB', '纯视野效应'), seg('bg_maj', '纯视野效应')
assert (f2(-fs['diff']), f2(fs['ci'][0]), f2(fs['ci'][1]), sci(fs['p'])) == ('43.54', '-45.80', '-41.40', '1.0 × 10⁻⁸⁰')
assert (f2(-fm['diff']), f2(fm['ci'][0]), f2(fm['ci'][1]), sci(fm['p'])) == ('22.29', '-24.55', '-20.18', '2.6 × 10⁻⁴⁵')
assert f2(-fb['diff']) == '2.03' and f'{abs(fbg["diff"]):.2f}' == '0.01'
nb, nbg = seg('tgtB', '纳入无结节切片'), seg('bg_maj', '纳入无结节切片')
assert (f2(-nb['diff']), f2(nb['ci'][0]), f2(nb['ci'][1])) == ('14.73', '-16.78', '-12.63')
assert (f2(nbg['diff']), f2(nbg['ci'][0]), f2(nbg['ci'][1])) == ('1.62', '0.78', '2.58')
assert (f4(lad('tgtBshift', 'win128')), f4(lad('tgtBshift', 'win512'))) == ('0.7598', '0.3243')
assert (f4(lad('tgtBshift_bg', 'win128')), f4(lad('tgtBshift_bg', 'win512'))) == ('0.7601', '0.5372')
assert (f4(lad('tgtB', 'patch128')), f4(lad('tgtB', 'win128')), f4(lad('bg_maj', 'slice'))) == ('0.8720', '0.2552', '0.0505')
t3 = SC['table3_per_scope_opt']
assert f2(100 * (t3['tgtBshift_v2_win128']['dice_opt'] - t3['tgtBshift_v2_win512']['dice_opt'])) == '29.67'
assert (t3['tgtBshift_v2_win128']['thr'], t3['tgtBshift_v2_win384']['thr'], t3['tgtBshift_v2_win512']['thr']) == (0.925, 0.998, 0.998)
pu = T['archPU1e3']
PU_NEW = [f"{f4(pu[k]['dice_05'])} / {f4(pu[k]['dice_opt'])}" for k in '1234']
assert PU_NEW == ['0.8218 / 0.8487', '0.8599 / 0.8601', '0.7444 / 0.7619', '0.6319 / 0.6751'], PU_NEW
pre = [T[a]['2']['dice_opt'] for a in ('tgtB', 'archR18', 'archCNX', 'archPVT')]
assert f2(100 * (max(pre) - min(pre))) == '0.28'
mass = J('outputs', 'analysis_test', 'symmetric_analysis_applied.json')['threshold_free']['prob_mass']
others = [v for k, v in mass.items() if k not in ('A Union (V>=1)', 'tgtBshift', 'tgtF')]
assert len(others) == 8 and (f'{min(others):.1f}', f'{max(others):.1f}') == ('183.2', '193.0'), others


# ----------------------------------------------------------------- helpers（与 v21 相同）
def find_para(doc, prefix, contains=None):
    hits = [p for p in doc.paragraphs
            if p.text.startswith(prefix) and (contains is None or contains in p.text)]
    assert len(hits) == 1, f'find_para({prefix[:50]!r}) -> {len(hits)} hits'
    return hits[0]


def _check_runs(p, tag):
    joined = ''.join(r.text for r in p.runs)
    assert joined == p.text, f'[{tag}] paragraph has text outside plain runs; refuse to edit'


def replace(p, old, new, tag):
    """在单个 run 内做替换（不跨 run，避免破坏斜体等格式）。"""
    _check_runs(p, tag)
    full = p.text
    assert full.count(old) == 1, f'[{tag}] expected 1 x {old[:60]!r}, found {full.count(old)}'
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            break
    else:
        raise AssertionError(f'[{tag}] {old[:60]!r} spans several runs')
    assert p.text == full.replace(old, new), f'[{tag}] post-check failed'
    LOG.append(tag)


def set_single_run_text(p, new, tag):
    _check_runs(p, tag)
    assert len(p.runs) == 1, f'[{tag}] expected a single run, found {len(p.runs)}'
    p.runs[0].text = new
    LOG.append(tag)


def png_size(path):
    b = Path(path).read_bytes()[:24]
    assert b[:8] == b'\x89PNG\r\n\x1a\n'
    return struct.unpack('>II', b[16:24])


# =================================================================== MAIN
doc = docx.Document(str(SRC['main']))

# ---- 1. 摘要 / Highlights / 贡献
set_single_run_text(find_para(doc, 'Results. On the held-out test set'),
    'Results. On the held-out test set, changing only the evaluation target moved Dice by 28.58 points and only the '
    f'threshold by 9.99; widening the evaluation window moved it by {N["fov"]} points after both training-side '
    f'confounds were removed, and an 8-pixel window offset cost a median of {N["off_med"]} points. The largest '
    f'learning-design effect was {N["sup"]} points, and the architecture range ({N["arch"]} points) did not exceed '
    'the 1.19-point run-to-run noise floor; even the conservative field-of-view estimate exceeded these by '
    f'{N["r_sup"]}× and {N["r_arch"]}×, respectively. Cross-validation reproduced this separation. Soft supervision '
    'over rater votes conferred no measurable benefit. On seven external multi-rater tasks, the protocol effect fell '
    'as rater agreement rose.', '1a abstract results (1.19, 22.29, 26.40, 7.1x, 18.7x)')
replace(find_para(doc, '• Evaluation field of view costs'), '22.5 Dice points', '22.3 Dice points', '1b highlight 22.3')
replace(find_para(doc, '3. A direct comparison of protocol effect sizes'),
        'by 7.2× and the architecture range by 12.0× on the conservative held-out-test estimate, with the same '
        'ordering reproduced in cross-validation.',
        f'by {N["r_sup"]}× and the architecture range by {N["r_arch"]}× on the conservative held-out-test estimate, '
        'with the same separation reproduced in cross-validation.', '1c contribution 3 ratios')

# ---- 2. §3.1 名义像素间距披露；Table 1 图题
p = find_para(doc, 'Nodules were defined as connected components of the majority-vote mask')
replace(p, 'and lesions endorsed by a single reader.',
        'and lesions endorsed by a single reader. Slices are used at their native in-plane resolution, and component '
        "areas are converted to millimetres at a nominal spacing of 0.6 mm per pixel rather than at each scan's own "
        'spacing; the inclusion criterion is therefore an area of at least 20 pixels, and all diameters and size '
        'strata in this paper are on this nominal scale.', '2a nominal 0.6 mm/px disclosure')
replace(find_para(doc, 'Table 1. Cohort composition.'),
        'Diameters are equivalent circular diameters of the majority-vote component on each slice.',
        'Diameters are equivalent circular diameters of the majority-vote component on each slice, at the nominal '
        'pixel spacing of 0.6 mm (Section 3.1).', '2b Table 1 caption nominal spacing')

# ---- 3. §3.2 面积；Fig. 1 图注
replace(find_para(doc, 'These definitions differ substantially in extent.'),
        'The test set shows the same pattern at slightly smaller scale (111.8, 76.5, 37.8 and 19.3 px per contributing case).',
        'The held-out test set shows the same nesting with a slightly smaller unanimous core (245.3, 192.3, 126.3 and '
        '70.9 pixels over its 2062 instances).', '3a test-set areas (old values came from a discarded first-instance-per-case computation)')
p = find_para(doc, 'Fig. 1 — Annotation consensus levels')
replace(p, 'on a representative LIDC nodule', 'on an isolated LIDC nodule', '3b Fig. 1 title')
replace(p, ' (case LIDC-IDRI-0717, equivalent diameter 5.8 mm, within the 5–10 mm band that holds 46.9% of the cohort). '
           '(a) Lung-windowed CT patch.',
        ' (case LIDC-IDRI-0717; equivalent diameter of the majority mask 4.9 mm at the nominal 0.6 mm pixel spacing, '
        'in the 3–5 mm band that holds 25.0% of cross-validation nodules). (a) Lung-windowed CT, a 48 × 48-pixel crop '
        'of the 128-pixel training patch.', '3c Fig. 1 diameter on the nominal scale used by Table 1')

# ---- 4. §4.1–4.3
replace(find_para(doc, 'How much of this is arithmetic.'),
        'their case-level mean areas are 260.3, 202.4, 142.8 and 89.3 pixels',
        'their case-level mean areas in cross-validation are 209.7, 163.4, 116.8 and 73.5 pixels',
        '4a case-level areas (old values were run-weighted row means)')
replace(find_para(doc, 'Fig. 2 — Dice versus binarisation threshold'),
        'The optima migrate from below 0.1 against the union reference to above 0.97 against the unanimous one',
        'The optima migrate from 0.2 or below against the union reference (0.002–0.005 for arms B–D) to 0.98 or '
        'above against the unanimous one', '4b Fig. 2 caption: arm A optimum vs G1 is 0.2')
replace(find_para(doc, 'Fig. 4 — Training-target'),
        " — Training-target × evaluation-target cross-evaluation matrix. Each cell is Dice at that arm's calibrated "
        'threshold for that evaluation target. The near-diagonal structure is the finding: no supervision target is '
        'best everywhere.',
        ' — Training-target × evaluation-target cross-evaluation matrix for the five supervision targets of Sections '
        "4.3 and 4.12. Each cell is Dice, in points, at that arm's own calibrated threshold for that evaluation "
        'reference; the outlined cell in each column is the best arm against that reference. No supervision target '
        'is best everywhere — the union arm is best against G₁, the SVLS arm against G₂ and the majority arm against '
        'G₃ and G₄ — and the spread within each column (at most 2.65 points) is small beside the spread across columns.',
        '4c Fig. 4 caption: drop "near-diagonal", state best cells')
replace(find_para(doc, 'Fig. 5 — Forest plot'),
        'the shaded band is the run-to-run noise floor. The bottom block (C − B) isolates soft supervision: its two '
        'upper intervals sit inside the noise band, which is the null result.',
        'the shaded band is the run-to-run noise floor (±1.19 points, 2σ). Grey labels mark intervals that include '
        'zero; these three are also the only contrasts in the figure that are not significant after Holm correction '
        '(Table 4). The bottom block (C − B) isolates soft supervision: its intervals against G₁ and G₂ lie inside '
        'the noise band and include zero, which is the null result.', '4d Fig. 5 caption: G1/G2 are the lower rows; noise floor value; Holm')
replace(find_para(doc, 'The cross-evaluation matrix has a near-diagonal structure'),
        'The cross-evaluation matrix has a near-diagonal structure (Fig. 4).',
        'No supervision target is best against every reference (Fig. 4).', '4e matrix wording')

# ---- 5. §4.5 效应量表与 Fig. 3
replace(find_para(doc, 'Table 5 places all measured effects on one axis'),
        'Table 5 reports the cross-validation estimates of Sections 4.1–4.6 together with the scope estimates of '
        'Sections 4.11 and 4.13; Fig. 3 shows the same ranking computed entirely on the held-out test set, so that '
        'the figure and the table are independent of one another.',
        'Table 5 reports the cross-validation estimates of Sections 4.1–4.6 together with the offset and scope '
        'estimates of Sections 4.10, 4.11 and 4.13, which require whole-slice inference and are therefore measured on '
        'the held-out set; Fig. 3 recomputes every remaining entry on the held-out test set as well. The two agree on '
        'every comparison between a protocol factor and a learning-design factor and differ only in the order of two '
        'protocol entries of similar size, the evaluation-target and offset effects (25.58 and 26.40 points in '
        'Table 5, 28.58 and 26.40 in Fig. 3).', '5a Table 5 vs Fig. 3: orders differ at ranks 2-3')
replace(find_para(doc, 'Table 5. Effect sizes measured in this study'),
        'Entries are cross-validation estimates except the two evaluation-field-of-view rows and the '
        'translation-augmentation cost on test, which require whole-slice inference and are therefore measured on the '
        '158-case held-out set; the source section is given for each.',
        'Entries are cross-validation estimates except the two evaluation-field-of-view rows and the offset row, '
        'which require whole-slice inference and are therefore measured on the held-out set (158 contributing cases; '
        '76 eligible nodules for the offset row), and the second value of the translation-augmentation row; the '
        'source section is given for each.', '5b Table 5 caption: offset row is a test-set measurement')
p = find_para(doc, 'Two field-of-view entries appear')
replace(p, 'The 43.84-point figure', f'The {N["fov_pi"]}-point figure', '5c 43.54')
replace(p, 'The 22.51-point figure', f'The {N["fov"]}-point figure', '5d 22.29')
set_single_run_text(find_para(doc, 'We use the conservative figure for every comparative claim'),
    'We use the conservative figure for every comparative claim in this paper. On that basis, on the held-out test '
    'set, the conservative field-of-view effect exceeds the largest learning-design effect (the supervision target) '
    f'by {N["r_sup"]}× and the architecture range by {N["r_arch"]}× — descriptive point estimates on the full test '
    'set (n = 158); on the 124-case common subset (non-empty V ≥ 4 reference across all arms) the joint bootstrap '
    'estimates are 7.0× [5.8, 8.7] and 15.6× [9.3, 30.9], with the same order-of-magnitude conclusion; on the '
    f'position-invariant arm the ratios are {N["r_pi_sup"]}× and {N["r_pi_arch"]}×. Under either convention, every '
    'evaluation-protocol factor in Table 5 exceeds every learning-design factor, as does the position prior '
    'installed by nodule-centred cropping.', '5e ratios (effect_uncertainty.json)')
replace(find_para(doc, 'On validation the architecture range is 0.8 times the noise floor'),
        'on the test set it widens to 1.87 points (Section 4.6).',
        'on the test set it widens to 1.19 points, level with the floor (Section 4.9).', '5f test architecture range')
p = find_para(doc, 'Two further entries deserve care.')
replace(p, 'it would appear to be 2.09 points', f'it would appear to be {f2(-fb["diff"])} points', '5g 2.03')
replace(p, 'collapsed to 0.2566 by the first window', f'collapsed to {f4(lad("tgtB", "win128"))} by the first window', '5h 0.2552')
replace(p, '(6.9 to 35.5 points)', f'({N["off_lo"]} to {N["off_hi"]} points)', '5i offset range')
p = find_para(doc, 'Fig. 3 — Effect sizes')
replace(p, ' = 158 cases). Bars are colour-coded by category; error bars are 95% case-level bootstrap confidence '
           'intervals where the quantity is a paired contrast. The shaded band at the left marks the run-to-run noise '
           'floor, which is the one quantity in the figure estimated on the validation folds (Section 4.8). The '
           'test-set values differ slightly from the cross-validation estimates of Table 5 — the evaluation-target '
           'effect is 28.58 rather than 25.58 points, the supervision-target range 3.12 rather than 2.65 and the '
           'architecture range 1.87 rather than 0.96, the threshold effect 9.99 rather than 9.42 — but the ordering is '
           'identical under both.',
        ' = 158 contributing cases; the offset entry on its 76 eligible nodules). Bars are colour-coded by category. '
        'Error bars are 95% case-level bootstrap confidence intervals — for the three entries defined as a range '
        '(evaluation target, supervision target, architecture), intervals of the range itself — and, for the noise '
        'floor, the 95% confidence interval of 2σ; the offset entry is the median '
        'over ten independently trained arms, whose individual values are shown as dots. The shaded band at the left '
        'marks the run-to-run noise floor, which is the one quantity in the figure estimated on the validation folds '
        '(Section 4.8). The test-set values differ slightly from the cross-validation estimates of Table 5 — the '
        'evaluation-target effect is 28.58 rather than 25.58 points, the threshold effect 9.99 rather than 9.42, the '
        'supervision-target range 3.12 rather than 2.65 and the architecture range 1.19 rather than 0.96 — so the '
        'evaluation-target and offset entries exchange places, but every protocol factor exceeds every learning-design '
        'factor under both.', '5j Fig. 3 caption (CIs on every bar but the median; order)')
set_single_run_text(find_para(doc, 'The test set is therefore the primary result'),
    'The test set is therefore the primary result: evaluation target 28.58 points, supervision target 3.12, '
    'architecture 1.19 (Section 4.9 and Fig. 3), with the cross-validation analysis of Sections 4.1–4.8 reproducing '
    'the same separation between protocol and learning-design effects.', '5k primary result')

# ---- 6. §4.7 概率质量（逐实例口径写明）；Fig. 6 图注
replace(find_para(doc, 'The arms differ in how they distribute probability mass.'),
        'Summed over the patch, arms B, C and D output mean probability masses of 192.9, 190.7 and 188.5, all close '
        'to the mean area of the majority reference (203.2 px).',
        'Summed over the patch and averaged over lesion instances, arms B, C and D output mean probability masses of '
        '192.9, 190.7 and 188.5, all close to the per-instance mean area of the majority reference (203.2 px).',
        '6a probability mass unit stated')
p = find_para(doc, 'Fig. 6 — (a) Mean predicted area')
replace(p, '(a) Mean predicted area as a function of the binarisation threshold, with the mean areas of ',
        '(a) Case-level mean predicted area as a function of the binarisation threshold in cross-validation, with the '
        'case-level mean areas of ', '6b Fig. 6 caption unit')
replace(p, '₂ marked. The union arm', '₂ marked (209.7 and 163.4 px); dots mark the area-matched operating points '
        'of Section 4.7. The union arm', '6c Fig. 6 caption reference values')

# ---- 7. §4.9 测试集
replace(find_para(doc, '• Protocol dominance.'), 'The architecture range at V ≥ 2 was 1.87 points.',
        'The architecture range at V ≥ 2 was 1.19 points, equal to the run-to-run noise floor.', '7a test architecture range')
set_single_run_text(find_para(doc, 'One quantity does not replicate at the same magnitude'),
    'One quantity is larger on the test set than on validation, and we report it as such: the architecture range '
    'grows from 0.96 points to 1.19. On both sets it is the gap between the plain U-Net and the best pretrained '
    'encoder, while the four pretrained encoders span only 0.19 and 0.28 points respectively, and on the test set it '
    'is level with the 1.19-point noise floor. The ordering evaluation protocol ≫ learning design survives with a '
    'wide margin. The supervision-target range (3.12 points, 95% CI [2.53, 3.71]) now clearly exceeds the '
    'architecture range (1.19, [0.57, 1.96]), but the two are measured against different references, and we do not '
    'rank them against each other on this evidence. What we claim is the part that is stable: both are at least '
    'three times smaller than every evaluation-protocol effect.', '7b architecture paragraph (no "order of magnitude")')
set_single_run_text(find_para(doc, 'The over-dispersion of the union arm also replicated'),
    'The over-dispersion of the union arm also replicated: its mean probability mass was 219.3 against 183.2–193.0 '
    'for the other eight arms of Table 8.', '7c probability mass range on test')

# ---- 8. §4.10 位置先验
p = find_para(doc, 'For arm B, averaged over its five fold models')
replace(p, '₂ falls from 0.8731 at ', f'₂ falls from {f4(cb["0"])} at ', '8a d0')
replace(p, ' = 0 to 0.5730 at ', f' = 0 to {f4(cb["8"])} at ', '8b d8')
replace(p, 'and to 0.0345 at ', f'and to {f4(cb["56"])} at ', '8c d56')
replace(p, 'is +60.72 points (95% CI [+55.40, +65.75], ',
        f'is {pf["mean_pts"]:+.2f} points (95% CI [{pf["ci95_pts"][0]:+.2f}, {pf["ci95_pts"][1]:+.2f}], ', '8d d40 paired')
replace(p, 'and a case-level cluster bootstrap returns the same interval, [+55.31, +65.85].',
        'and a case-level cluster bootstrap is therefore identical to the nodule-level one.', '8e cluster bootstrap')
p = find_para(doc, 'The effect is not specific to one architecture or one supervision target.')
replace(p, 'lose between 6.9 and 35.5 points at ', f'lose between {N["off_lo"]} and {N["off_hi"]} points at ', '8f range d8')
replace(p, ' = 8 (median 26.36) and between 66.9 and 86.4 points at ',
        f' = 8 (median {N["off_med"]}) and between {N["off56_lo"]} and {N["off56_hi"]} points at ', '8g median, range d56')
replace(find_para(doc, 'Fig. 7 — Dice against'),
        'The shaded region marks the ±32-pixel range used during augmentation;',
        'Each arm is scored at its own validation-calibrated threshold. The shaded region marks the ±32-pixel range '
        'used during augmentation;', '8h Fig. 7 caption threshold rule')

# ---- 9. §4.11 视野阶梯
replace(find_para(doc, 'The spatial extent over which Dice is computed'),
        'is the protocol factor least often reported and, we find, the largest.',
        'is the protocol factor least often reported and, on a conventionally trained position-invariant model, the '
        'largest we measured.', '9a "the largest" qualified')
replace(find_para(doc, 'One implementation point decides how this experiment may be read.'),
        'and is reported separately.',
        'and is reported separately. As in Section 4.10, each arm is scored at its own validation-calibrated threshold.',
        '9b ladder threshold rule stated')
p = find_para(doc, 'The first result is methodological.')
replace(p, '0.8719 at the 128-pixel input crop to 0.2566',
        f'{f4(lad("tgtB", "patch128"))} at the 128-pixel input crop to {f4(lad("tgtB", "win128"))}', '9c arm B ladder')
replace(p, 'declines by only 2.09 points', f'declines by only {f2(-fb["diff"])} points', '9d 2.03')
replace(p, 'that 2.09 says', f'that {f2(-fb["diff"])} says', '9e 2.03')
replace(p, 'sits at 0.0511 and moves', f'sits at {f4(lad("bg_maj", "slice"))} and moves', '9f 0.0505')
replace(find_para(doc, 'On the translation-augmented arm, which does not collapse'),
        '₂ falls from 0.7590 at the 128-pixel window to 0.3206 at 512 pixels: 43.84 points (95% CI [−46.09, −41.69], '
        'p = 2.4 × 10⁻⁸¹, paired by case, ',
        f'₂ falls from {f4(lad("tgtBshift", "win128"))} at the 128-pixel window to {f4(lad("tgtBshift", "win512"))} '
        f'at 512 pixels: {N["fov_pi"]} points (95% CI [−45.80, −41.40], p = {sci(fs["p"])}, paired by case, ',
        '9g translation-augmented ladder')
set_single_run_text(find_para(doc, 'Recalibrating the threshold at each scope'),
    'Recalibrating the threshold at each scope recovers only part of the loss — scored at each window\'s own optimal '
    'threshold, the translation-augmented arm still loses 29.67 points between the 128- and 512-pixel windows — and '
    'the optimal threshold rises with window size, from 0.925 at 128 pixels to 0.998 at 384 and 512 pixels. That '
    'drift is itself a reportable regularity, and a further reason why a threshold selected under one evaluation '
    'extent cannot be transferred to a study using another.', '9h per-scope thresholds (was "0.95 to 0.998")')
replace(find_para(doc, 'Fig. 8 — (a) Dice against'),
        'The shaded band marks the field-of-view ladder proper.',
        'The shaded band marks the field-of-view ladder proper; every arm is scored at its own validation-calibrated '
        'threshold.', '9i Fig. 8 caption threshold rule')

# ---- 10. §4.13
p = find_para(doc, 'The background-sampled arm establishes')
replace(p, 'by a further 14.84 points (95% CI [−16.91, −12.73], paired by case, ',
        'by a further 14.73 points (95% CI [−16.78, −12.63], paired by case, ', '10a nodule-free step, arm B')
replace(p, 'the same step is +1.60 points ([+0.76, +2.56])', 'the same step is +1.62 points ([+0.78, +2.58])', '10b background arm')
replace(p, 'the background arm sits at 0.0511 on full-slice input', 'the background arm sits at 0.0505 on full-slice input', '10c 0.0505')
replace(find_para(doc, 'The merged arm does.'),
        "it holds 0.7600 at the 128-pixel window, comparable to the translation-augmented arm's 0.7590 — and across "
        'the ladder to 512 pixels it loses 22.51 points (95% CI [−24.77, −20.41], p = 9.9 × 10⁻⁴⁶, paired by case, ',
        f"it holds {f4(lad('tgtBshift_bg', 'win128'))} at the 128-pixel window, comparable to the translation-augmented "
        f"arm's {f4(lad('tgtBshift', 'win128'))} — and across the ladder to 512 pixels it loses {N['fov']} points "
        f'(95% CI [−24.55, −20.18], p = {sci(fm["p"])}, paired by case, ', '10d merged arm')
replace(find_para(doc, 'First, about half of the'), 'the 43.84-point figure', f'the {N["fov_pi"]}-point figure', '10e 43.54')
set_single_run_text(find_para(doc, 'Second, the other half is not.'),
    f'Second, the other half is not. {N["fov"]} points survive on a model for which we can identify no remaining '
    f'training-side confound. That is still {N["r_arch"]} times the run-to-run noise floor — and, because the '
    f'test-set architecture range equals that floor, {N["r_arch"]} times the architecture range — and {N["r_sup"]} '
    'times the largest supervision-target effect. The evaluation field of view therefore remains one of the largest '
    'factors measured in this study, and the conclusion of Section 4.5, that every protocol factor exceeds every '
    'learning-design factor, is unchanged in direction, only in magnitude.', '10f "largest single factor" corrected')
p = find_para(doc, 'Third, the two remedies are not interchangeable')
replace(p, 'leaves 43.84 points on the ladder', f'leaves {N["fov_pi"]} points on the ladder', '10g')
replace(p, 'both together leave 22.51.', f'both together leave {N["fov"]}.', '10h')
replace(p, 'we do not claim that 22.51 is irreducible', f'we do not claim that {N["fov"]} is irreducible', '10i')

# ---- 11. 讨论 / 局限 / 结论
replace(find_para(doc, 'The central observation is arithmetic.'), 'at 0.7600 or 0.5349',
        f'at {f4(lad("tgtBshift_bg", "win128"))} or {f4(lad("tgtBshift_bg", "win512"))}', '11a merged-arm ladder ends')
replace(find_para(doc, 'Several published differences cited in this paper'), '22.51 for the field of view',
        f'{N["fov"]} for the field of view', '11b')
p = find_para(doc, '3. The evaluation extent:')
replace(p, 'Measured effect: 22.51 points', f'Measured effect: {N["fov"]} points', '11c')
replace(p, '43.84 points on a conventionally trained one', f'{N["fov_pi"]} points on a conventionally trained one', '11d')
replace(find_para(doc, '4. The patch-cropping rule at training time'), 'a median of 26.36 points',
        f'a median of {N["off_med"]} points', '11e')
replace(find_para(doc, 'Architecture family.'), 'The range also widens to 1.87 points on the held-out test set',
        'The range also widens to 1.19 points on the held-out test set', '11f')
p = find_para(doc, 'The residual field-of-view effect may not be irreducible.')
replace(p, 'Section 4.13 reports 22.51 points', f'Section 4.13 reports {N["fov"]} points', '11g')
replace(p, 'which is why 22.51 is presented', f'which is why {N["fov"]} is presented', '11h')
p = find_para(doc, 'Widening the evaluation window from 128 to 512 pixels')
replace(p, 'costs 22.51 Dice points', f'costs {N["fov"]} Dice points', '11i')
replace(p, 'the same manipulation costs 43.84.', f'the same manipulation costs {N["fov_pi"]}.', '11j')
replace(p, 'costs a median of 26.36 points', f'costs a median of {N["off_med"]} points', '11k')
set_single_run_text(find_para(doc, 'Against these, the largest difference among five defensible supervision'),
    'Against these, the largest difference among five defensible supervision strategies is 3.12 points, and the '
    'range across five architectures spanning convolutional and transformer families and an 8-fold parameter range '
    'is 1.19 points — no larger than the 1.19-point noise floor obtained by retraining with different seeds — of '
    'which the four ImageNet-pretrained encoders account for 0.28. On the conservative estimate, evaluation-protocol '
    f'choices exceed learning-design choices by {N["r_sup"]}× and the architecture range by {N["r_arch"]}×; '
    'cross-validation, reported as a development analysis, reproduces the same separation.', '11l conclusions ratios')
replace(find_para(doc, 'During the preparation of this work the authors used'),
        'Claude (Anthropic), Doubao (ByteDance) and Codex (OpenAI)',
        'Claude (Anthropic), Doubao (ByteDance), ChatGLM (Zhipu AI) and Codex (OpenAI)', '11m AI declaration: add Zhipu')

# ---- 12. 表格
def cell_replace(cell, old, new, tag):
    hits = [p for p in cell.paragraphs if old in p.text]
    assert len(hits) == 1, f'[{tag}] {len(hits)} hits'
    replace(hits[0], old, new, tag)


t5 = doc.tables[4]
assert t5.rows[0].cells[0].text.startswith('Factor')
cell_replace(t5.rows[1].cells[2], '43.84', N['fov_pi'], '12a Table 5 FOV (position-invariant)')
cell_replace(t5.rows[2].cells[2], '26.36 (median of 10 arms; range 6.9–35.5)',
             f'{N["off_med"]} (median of 10 arms; range {N["off_lo"]}–{N["off_hi"]})', '12b Table 5 offset')
cell_replace(t5.rows[4].cells[2], '22.51', N['fov'], '12c Table 5 FOV (merged)')
vals = [float(re.match(r'([\d.]+)', r.cells[2].text.strip()).group(1)) for r in t5.rows[1:] if r.cells[1].text != 'Noise floor']
assert vals == sorted(vals, reverse=True), vals
t8 = doc.tables[7]
row = [r for r in t8.rows if r.cells[0].text.startswith('Plain U-Net')]
assert len(row) == 1
for c, old, new in zip(row[0].cells[1:5], ['0.8121 / 0.8448', '0.8532 / 0.8532', '0.7438 / 0.7583', '0.6326 / 0.6774'], PU_NEW):
    cell_replace(c, old, new, f'12d Table 8 Plain U-Net {old} -> {new}')

# ---- 13. 插图：按图注编号定位每张图，换成 outputs/figures/v22/FigN.png，宽度取图像自然宽度（≤16.51 cm）
paras = doc.paragraphs
mapping = {}
for i, p in enumerate(paras):
    blips = p._p.findall('.//' + qn('a:blip'))
    if not blips:
        continue
    assert len(blips) == 1, f'paragraph {i} holds {len(blips)} images'
    cap = next((q for q in paras[i + 1:i + 4] if re.match(r'^Fig\. \d+ —', q.text)), None)
    assert cap is not None, f'no caption after image paragraph {i}'
    n = int(re.match(r'^Fig\. (\d+)', cap.text).group(1))
    assert n not in mapping
    mapping[n] = (p, blips[0].get(qn('r:embed')))
assert sorted(mapping) == list(range(1, 10)), sorted(mapping)
EMU_IN, MAXW = 914400, int(16.51 / 2.54 * 914400)
for n, (p, rid) in sorted(mapping.items()):
    png = FIGS / f'Fig{n}.png'
    w, h = png_size(png)
    part = doc.part.related_parts[rid]
    part._blob = png.read_bytes()
    cx, cy = int(w / 600 * EMU_IN), int(h / 600 * EMU_IN)
    if cx > MAXW:
        cy, cx = int(cy * MAXW / cx), MAXW
    for tag in ('wp:extent', 'a:ext'):
        for el in p._p.iter(qn(tag)):
            el.set('cx', str(cx)); el.set('cy', str(cy))
    LOG.append(f'13 Fig. {n} <- {png.name} ({w}x{h}px, {cx / 360000:.2f} x {cy / 360000:.2f} cm)')

# ---- 14. 残留检查：旧数字不得再出现（正文段落 + 表格）
STALE = ['22.51', '43.84', '26.36', '1.87', '12.0×', '0.7590', '0.3206', '0.5349', '0.2566', '0.0511', '14.84',
         '60.72', '55.40', '55.31', '0.8731', '0.5730', '0.0345', '66.9', '86.4', '260.3', '111.8', '5.8 mm',
         'below 0.1', 'near-diagonal', '183.2–188.7', 'twenty times', 'two upper intervals', '0.95 to 0.998',
         '10.4×', '23.4×', 'an order of magnitude smaller than the evaluation-protocol', 'largest single factor',
         '0.8121', '0.8532', '7.2×', '7.2 times']
texts = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
left = [(s, t[:80]) for s in STALE for t in texts if s in t]
assert not left, left

DST['main'].parent.mkdir(parents=True, exist_ok=True)
doc.save(str(DST['main']))

# =================================================================== SUPPLEMENT：删除孤立图片
supp = docx.Document(str(SRC['supp']))
used = {b.get(qn('r:embed')) for b in supp.element.body.iter(qn('a:blip'))}
dropped = [rid for rid, rel in list(supp.part.rels.items()) if rel.reltype.endswith('/image') and rid not in used]
for rid in dropped:
    supp.part.drop_rel(rid)
LOG.append(f'S1 supplement: dropped {len(dropped)} unreferenced image parts (old figure copies)')
supp.save(str(DST['supp']))

# =================================================================== HIGHLIGHTS / COVER LETTER
hl = docx.Document(str(SRC['hl']))
replace(find_para(hl, 'Evaluation field of view costs'), '22.5 Dice points', '22.3 Dice points', 'H1 highlights 22.3')
for p in hl.paragraphs[1:]:
    assert len(p.text) <= 85, p.text
hl.save(str(DST['hl']))

cv = docx.Document(str(SRC['cover']))
p = find_para(cv, 'On a held-out test set of 202 cases')
replace(p, 'moved it by 22.51 points', f'moved it by {N["fov"]} points', 'L1 cover 22.29')
replace(p, 'cost a median of 26.36 points', f'cost a median of {N["off_med"]} points', 'L2 cover 26.40')
hl_run = [r for r in p.runs if r.text == '1.87 points']      # v21 里这一处是黄色高亮的待改标记
assert len(hl_run) == 1
hl_run[0].text = '1.19 points'
hl_run[0].font.highlight_color = None
replace(p, ', against a 1.19-point run-to-run noise floor', ', no larger than the 1.19-point run-to-run noise floor',
        'L3 cover architecture (highlight removed: resolved)')
cv.save(str(DST['cover']))

print(f'{len(LOG)} edits:')
for t in LOG:
    print('  ·', t)
for k, v in DST.items():
    print(f'  -> {v}')
