# -*- coding: utf-8 -*-
"""revise_paper_v18.py —— v17 -> v18：QUBIQ 8 点改写之后遗留的矛盾与错误。

复核 v17 时查出 7 处（详见 docs/v17_复核_20260905.md）：
  1. Highlight 第 5 条仍宣称"比值随一致性单调下降"——该趋势已被证伪
     （8 点 rho=-0.31 p=0.46；7 点 rho=-0.07 p=0.91）
  2. §5.5 仍写"三个数据集 / 一致性指数 0.33, 0.49, 0.87"，与 §4.14 直接矛盾
  3. §6 结论仍写"across three multi-rater datasets it falls monotonically"
  4. Fig. 9 图注把 rho=-0.86 和 p=0.048 配在一起（后者是 7 点的 p）
  5. §4.14 的 p = 0.007 是 t 近似；精确置换（8! = 40320）是 p = 0.011，
     而同一句把 7 点的 p 标为 "exact permutation"，两者方法不一致
  6. Table 7 表头 "Agreement inde×"（比值列的 x->× 全局替换误伤）
  7. §5.1 的 "(Table 6)" 指错表——学习率那一行在无编号的噪声底表里，
     Table 6 是保留测试集表。（这一处是 v17 那一轮我自己写错的）
另加一处补充：
  8. §3.8 说明一致性指数是在实际参与评测的 2D 切片上计算，不是原始 3D 体
"""
import json
import os

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = os.path.join(ROOT, 'outputs', 'paper')
SRC = os.path.join(P, '论文_协议效应_v17_CIBM.docx')
DST = os.path.join(P, '论文_协议效应_v18_CIBM.docx')

REPLACE = [
    # 1. Highlight 第 5 条
    ('• Protocol dominance falls monotonically as inter-rater agreement rises.',
     '• Protocol effects shrink as rater agreement rises across eight multi-rater tasks.'),
    # 4. Fig. 9 图注
    ('(Spearman rho = -0.86, permutation p = 0.048)',
     '(Spearman rho = -0.86, exact permutation p = 0.011)'),
    # 5. §4.14 的 p
    ('including LIDC as a reference point yields rho = -0.86 (p = 0.007)',
     'including LIDC as a reference point yields rho = -0.86 (exact permutation p = 0.011)'),
    # 7. §5.1 指错表
    ('the loss composition (Supplementary Note S1) and the encoder learning rate (Table 6)',
     'the loss composition (Supplementary Note S1) and the encoder learning rate (Section 4.8)'),
    # 2. §5.5 混杂说明
    ('Rater count is confounded with agreement. With three datasets, the rater count (3, 4, 7) '
     'and the agreement index (0.33, 0.49, 0.87) cannot be separated. The monotone relationship '
     'of Section 4.14 is an observed association, not an attributed mechanism.',
     'Rater count is only partly separated from agreement. Across the eight points of Section '
     '4.14 the rater counts are 3, 3, 3, 3, 4, 6, 6 and 7, so the count is no longer collinear '
     'with the agreement index as it was with three datasets alone; but the three brain-tumour '
     'subtasks share readers and images, as do the two prostate subtasks, so the eight points '
     'are not fully independent. The relationship of Section 4.14 is an observed association, '
     'not an attributed mechanism.'),
    # 3. §6 结论
    ('Finally, the dominance of protocol over method is not universal but graded: across three '
     'multi-rater datasets it falls monotonically as rater agreement rises, and on a task where '
     'three raters agree closely it disappears.',
     'Finally, the dominance of protocol over method is not universal but graded: across eight '
     'multi-rater tasks spanning four organs and two modalities, the absolute protocol effect '
     'falls as rater agreement rises (Spearman rho = −0.86, exact permutation p = 0.011), and on '
     'tasks where the raters agree closely it becomes small in absolute terms.'),
    # 9. 摘要末句：'the ordering ... falling with agreement' 读起来像比值趋势，
    #    改为明确指向绝对协议效应（词数 +2，仍 <= 250）
    ('and the ordering replicated externally, falling with inter-rater agreement.',
     'and the ordering replicated externally, the protocol effect falling with agreement.'),
    # 8. §3.8 口径说明
    ('It is 0 when no pixel is unanimous and 1 when the raters agree exactly.',
     'It is 0 when no pixel is unanimous and 1 when the raters agree exactly. The index is '
     'computed on the same two-dimensional slices that enter training and evaluation rather than '
     'on the full three-dimensional volumes, so that it describes the annotations the models '
     'actually see; on volumetric tasks the two can differ appreciably.'),
]


def run_replace(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    joined = ''.join(r.text for r in par.runs)
    if old not in joined:
        return False
    par.runs[0].text = joined.replace(old, new)
    for r in par.runs[1:]:
        r.text = ''
    return True


def main():
    doc = docx.Document(SRC)
    for old, new in REPLACE:
        hit = any(run_replace(p, old, new) for p in doc.paragraphs)
        assert hit, f'未找到: {old[:70]!r}'

    # 6. Table 7 表头
    tb = doc.tables[8]
    c = tb.rows[0].cells[2]
    assert c.text.strip() == 'Agreement inde×', f'表头不是预期值: {c.text!r}'
    c.paragraphs[0].runs[0].text = 'Agreement index'
    for r in c.paragraphs[0].runs[1:]:
        r.text = ''

    doc.save(DST)

    d2 = docx.Document(DST)
    txt = '\n'.join(p.text for p in d2.paragraphs)
    for bad in ('falls monotonically as inter-rater agreement rises', 'p = 0.007',
                'externally, falling with inter-rater agreement',
                'With three datasets', 'across three multi-rater datasets',
                'permutation p = 0.048)', 'encoder learning rate (Table 6)'):
        assert bad not in txt, f'v18 仍含: {bad}'
    assert d2.tables[8].rows[0].cells[2].text.strip() == 'Agreement index', '表头未修好'
    for good in ('exact permutation p = 0.011', 'shrink as rater agreement rises',
                 '3, 3, 3, 3, 4, 6, 6 and 7', 'eight multi-rater tasks spanning four organs',
                 'two-dimensional slices that enter training'):
        assert good in txt, f'v18 缺少: {good}'
    hl = [p.text.strip().lstrip('•').strip() for p in d2.paragraphs
          if p.text.strip().startswith('•')][:5]
    print('Highlights 字符数:', [len(h) for h in hl], '(上限 85)')
    assert max(len(h) for h in hl) <= 85
    a = [i for i, p in enumerate(d2.paragraphs) if p.text.strip() == 'Abstract'][0]
    h = [i for i, p in enumerate(d2.paragraphs) if p.text.strip() == 'Highlights'][0]
    n = sum(len(p.text.split()) for p in d2.paragraphs[a + 1:h]
            if p.text.strip() and not p.text.strip().startswith('Keywords'))
    print(f'摘要 {n} 词；段落 {len(d2.paragraphs)}  表 {len(d2.tables)}  图 {len(d2.inline_shapes)}  '
          f"PENDING {txt.count('PENDING')}")
    print('回读自检通过 ->', DST)


if __name__ == '__main__':
    main()
