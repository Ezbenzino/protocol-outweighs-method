# -*- coding: utf-8 -*-
"""v19 补两张无编号表的编号 + 标题（方案 A 重排）：按文档顺序 Table 1-9。
- TABLE[1]（§3.3 监督靶设计）-> Table 2；TABLE[6]（§4.8 噪声底）-> Table 7
- 原 Table 2->3, 3->4, 4->5, 5->6, 6->8, 7->9
- 正文本文 Table 引用同步更新（10 处），Zhi 的 "Their Table" 不动
- 正文 §3.3/§4.8 各加一处对新增表的引用
"""
import shutil
import docx
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = str(ROOT / 'outputs' / 'paper' / '论文_协议效应_v19_CIBM.docx')
BAK = str(ROOT / 'outputs' / 'paper' / '论文_协议效应_v19_CIBM_backup_renum.docx')
shutil.copyfile(SRC, BAK)
print('已备份 v19 ->', BAK)

d = docx.Document(SRC)

def find_para(sub, must=True):
    hits = [i for i, p in enumerate(d.paragraphs) if sub in p.text]
    if len(hits) == 1:
        return hits[0]
    if must:
        raise SystemExit('定位失败(命中%d): %r' % (len(hits), sub))
    return None

def replace_in_para(para, old, new, count_all=False):
    """优先在单个 run 内替换；跨 run 则重建段落文本。返回替换次数。"""
    n = 0
    for r in para.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            n += 1
            if not count_all:
                return n
    if n == 0 and old in para.text:
        full = para.text
        for r in list(para.runs):
            r._element.getparent().remove(r._element)
        para.add_run(full.replace(old, new))
        n += 1
    return n

def update_caption_prefix(sub, new_prefix):
    """caption 段落第一个 run 是 'Table N.'（加粗），改为新编号。"""
    idx = find_para(sub)
    p = d.paragraphs[idx]
    assert p.runs and p.runs[0].text.startswith('Table '), 'caption 首 run 不是 Table: %r' % p.runs[0].text
    p.runs[0].text = new_prefix + '.'

def insert_caption_before_table(table, bold_prefix, body_text):
    p_el = OxmlElement('w:p')
    table._tbl.addprevious(p_el)
    p = Paragraph(p_el, table._parent)
    r1 = p.add_run(bold_prefix)
    r1.bold = True
    p.add_run(body_text)
    return p

# ============ 1. caption 编号更新（6 处，改首 run） ============
update_caption_prefix('Table 2. Cross-evaluation matrix', 'Table 3')
update_caption_prefix('Table 3. Case-level paired comparisons', 'Table 4')
update_caption_prefix('Table 4. Effect sizes measured', 'Table 5')
update_caption_prefix('Table 5. Architecture control', 'Table 6')
update_caption_prefix('Table 6. Held-out test set', 'Table 8')
update_caption_prefix('Table 7. Protocol (evaluation-target)', 'Table 9')
print('caption 编号更新：6 处')

# ============ 2. 正文本文 Table 引用更新（10 处） ============
repls = [
    ('contrasts of Table 3 are', 'contrasts of Table 4 are'),
    ('under the same conditions (Table 2)', 'under the same conditions (Table 3)'),
    ('on the full set (Table 2)', 'on the full set (Table 3)'),
    ('Table 4 places', 'Table 5 places'),
    ('Table 4 reports', 'Table 5 reports'),
    ('factor in Table 4 exceeds', 'factor in Table 5 exceeds'),
    ('excluded from Table 4', 'excluded from Table 5'),
    ('estimates of Table 4', 'estimates of Table 5'),
    ('reproduce on the test set (Table 6)', 'reproduce on the test set (Table 8)'),
    ('Table 7, Fig. 9', 'Table 9, Fig. 9'),
]
done = 0
for old, new in repls:
    idx = find_para(old)
    n = replace_in_para(d.paragraphs[idx], old, new)
    if n == 0:
        raise SystemExit('替换失败: %r' % old)
    done += n
print('正文引用更新：%d 处' % done)

# ============ 3. 插入 2 个新 caption ============
insert_caption_before_table(
    d.tables[1],
    'Table 2.',
    ' The four supervision targets. All arms use the identical network, optimiser, schedule, '
    'augmentation and data splits; they differ only in which signal defines the Dice and '
    'cross-entropy terms. G\u2081\u2013G\u2084 are the nested consensus references V \u2265 1 to '
    'V \u2265 4; p is the per-pixel mean of the four rater masks.')
insert_caption_before_table(
    d.tables[6],
    'Table 7.',
    ' Measured effects against the run-to-run noise floor (2\u03c3 = 1.19 Dice points at fixed '
    'threshold 0.5, pooled across evaluation targets; Section 4.8).')
print('插入 2 个新 caption')

# ============ 4. 正文加引用 ============
replace_in_para(d.paragraphs[find_para('they differ only in which signal the Dice and cross-entropy terms are fitted to.')],
                'fitted to.', 'fitted to (Table 2).')
replace_in_para(d.paragraphs[find_para('evaluated them at the thresholds already selected on the five validation folds.')],
                'validation folds.', 'validation folds (Table 7).')
print('正文加 2 处表引用')

d.save(SRC)
print('补编号完成')
