# -*- coding: utf-8 -*-
"""revise_paper_v16_tables.py —— v15 -> v16 第 2 遍：表格与插图改为病例级口径。

第 1 遍（revise_paper_v16_paras.py）已改正文数字；本遍重建 Table 2/3/4/5/6/7/8
的数值单元格，并替换随口径变化的六张插图（Fig. 2/3/4/5/6/9）。
Fig. 1/7/8 不变（分别是病例插图、结节级位移、病例级视野阶梯）。

所有数值从权威 JSON 现取，不手抄。
"""
import copy
import json
import os

import docx
from docx.shared import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = os.path.join(ROOT, 'outputs', 'paper')
SRC = os.path.join(P, '论文_协议效应_v16_CIBM.docx')
FIG = os.path.join(ROOT, 'outputs', 'figures')

V = json.load(open(os.path.join(ROOT, 'outputs/analysis_full/symmetric_analysis.json'), encoding='utf-8'))
T = json.load(open(os.path.join(ROOT, 'outputs/analysis_test/symmetric_analysis_applied.json'), encoding='utf-8'))
N = json.load(open(os.path.join(ROOT, 'outputs/analysis_full/noise_floor.json'), encoding='utf-8'))
G = json.load(open(os.path.join(ROOT, 'outputs/analysis/generalisation_table.json'), encoding='utf-8'))
DIS = json.load(open(os.path.join(ROOT, 'outputs/analysis/disagreement_index.json'), encoding='utf-8'))
t1v, t1t, t2v = V['table1'], T['table1'], V['table2']

FLOOR = N['fixed_0_5']['two_sigma_pts']


def setcell(cell, text):
    """替换单元格文本，保留首个 run 的格式。"""
    p = cell.paragraphs[0]
    if p.runs:
        rpr = p.runs[0]._element.rPr
        for r in list(p.runs)[1:]:
            r._element.getparent().remove(r._element)
        p.runs[0].text = text
        if rpr is not None and p.runs[0]._element.rPr is None:
            p.runs[0]._element.insert(0, copy.deepcopy(rpr))
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def thr_str(t):
    s = f'{t:.3f}'.rstrip('0').rstrip('.')
    return s if s else '0'


def sci(p):
    if p >= 1e-3:
        return f'{p:.2f}'
    m, e = f'{p:.1e}'.split('e')
    return f'{m}e−{int(e[1:]):02d}' if e[0] == '-' else f'{m}e+{int(e[1:]):02d}'


def holm(pvals):
    order = sorted(range(len(pvals)), key=lambda i: pvals[i])
    m = len(pvals)
    adj = [0.0] * m
    prev = 0.0
    for rank, i in enumerate(order):
        val = min(1.0, (m - rank) * pvals[i])
        prev = max(prev, val)
        adj[i] = prev
    return adj


def main():
    doc = docx.Document(SRC)
    tb = doc.tables

    # ---------- Table 2：验证集交叉矩阵（A–D）----------
    for row, arm in zip(tb[2].rows[1:], ['tgtA', 'tgtB', 'tgtC', 'tgtD']):
        for j, k in enumerate([1, 2, 3, 4], start=1):
            e = t1v[arm][str(k)]
            setcell(row.cells[j],
                    f"{e['dice_05']:.4f} / {e['dice_opt']:.4f} ({thr_str(e['thr'])})")

    # ---------- Table 3：配对对比 + Holm（20 个预设对比上校正）----------
    PAIRS = [('tgtC-tgtA', 4), ('tgtB-tgtA', 4), ('tgtC-tgtB', 4), ('tgtD-tgtC', 4), ('tgtD-tgtB', 4)]
    allp, key = [], []
    for pr, _ in PAIRS:
        for k in (1, 2, 3, 4):
            allp.append(t2v[pr][str(k)]['p_t']); key.append((pr, k))
    adj = dict(zip(key, holm(allp)))
    SHOW = ['tgtC-tgtA', 'tgtB-tgtA', 'tgtC-tgtB', 'tgtD-tgtB']
    ri = 1
    for pr in SHOW:
        for k in (1, 2, 3, 4):
            e = t2v[pr][str(k)]
            c = tb[3].rows[ri].cells
            setcell(c[2], str(e['n']))
            setcell(c[3], f"{100*e['diff']:+.2f}".replace('-', '−'))
            setcell(c[4], f"[{100*e['ci'][0]:+.2f}, {100*e['ci'][1]:+.2f}]".replace('-', '−'))
            setcell(c[5], sci(e['p_t']))
            setcell(c[6], f"{e['dz']:.3f}".replace('-', '−'))
            setcell(c[7], sci(adj[(pr, k)]))
            ri += 1

    # ---------- Table 4：效应量清单 ----------
    def span(src, arm, key_, ks=(1, 2, 3, 4)):
        v = [src[arm][str(k)][key_] for k in ks]
        return 100 * (max(v) - min(v))
    SUP = ['tgtA', 'tgtB', 'tgtC', 'tgtD', 'tgtE']
    ARCH = ['tgtB', 'archR18', 'archCNX', 'archPVT', 'archPU1e3']
    sup4 = 100 * (max(t1v[a]['4']['dice_opt'] for a in SUP) - min(t1v[a]['4']['dice_opt'] for a in SUP))
    arch2 = 100 * (max(t1v[a]['2']['dice_opt'] for a in ARCH) - min(t1v[a]['2']['dice_opt'] for a in ARCH))
    thr_eff = 100 * (t1v['tgtA']['4']['dice_opt'] - t1v['tgtA']['4']['dice_05'])
    shift_v = 100 * (t1v['tgtB']['2']['dice_opt'] - t1v['tgtBshift']['2']['dice_opt'])
    shift_t = 100 * (t1t['tgtB']['2']['dice_opt'] - t1t['tgtBshift']['2']['dice_opt'])
    T4 = {
        'Evaluation target, V ≥ 1 to V ≥ 4, threshold fixed at 0.5': f'{span(t1v, "tgtA", "dice_05"):.2f}',
        'Binarisation threshold, arm A against G₄': f'{thr_eff:.2f}',
        'Cost of translation augmentation on the centred protocol': f'{shift_v:.2f} (val) / {shift_t:.2f} (test)',
        'Supervision target, range across 5 arms A–E against G₄': f'{sup4:.2f}',
        'Architecture, range across 5 models, 8-fold parameter range': f'{arch2:.2f}',
        'Retraining with a different random seed (2σ)':
            f"{FLOOR} (95% CI [{N['fixed05_two_sigma_ci'][0]}, {N['fixed05_two_sigma_ci'][1]}])",
    }
    for row in tb[4].rows[1:]:
        lab = row.cells[0].text.strip()
        if lab in T4:
            setcell(row.cells[2], T4[lab])

    # ---------- Table 5：架构对照 ----------
    A5 = ['archPU1e3', 'archR18', 'tgtB', 'archCNX', 'archPVT']
    for row, arm in zip(tb[5].rows[1:6], A5):
        for j, k in enumerate([1, 2, 3, 4], start=4):
            setcell(row.cells[j], f"{t1v[arm][str(k)]['dice_opt']:.4f}")
    for j, k in enumerate([1, 2, 3, 4], start=4):
        setcell(tb[5].rows[6].cells[j], f'{span(t1v, A5[0], "dice_opt", (k,)) if False else 100*(max(t1v[a][str(k)]["dice_opt"] for a in A5)-min(t1v[a][str(k)]["dice_opt"] for a in A5)):.2f}')

    # ---------- Table 6：与噪声底比较 ----------
    et = span(t1v, 'tgtA', 'dice_05')
    sup2 = 100 * (max(t1v[a]['2']['dice_opt'] for a in SUP) - min(t1v[a]['2']['dice_opt'] for a in SUP))
    r18_34 = abs(100 * (t1v['archR18']['2']['dice_opt'] - t1v['tgtB']['2']['dice_opt']))
    cb2 = 100 * t2v['tgtC-tgtB']['2']['diff']
    cb4 = abs(100 * t2v['tgtC-tgtB']['4']['diff'])
    T6 = {
        'Evaluation-target spread, arm A at threshold 0.5': (f'{et:.2f}', f'yes, {et/FLOOR:.0f}×'),
        'Threshold, arm A against G₄': (f'{thr_eff:.2f}', 'yes'),
        'Training-target range against G₄': (f'{sup4:.2f}', 'yes'),
        'C − B (soft supervision) against G₄': (f'{cb4:.2f}', 'yes'),
        'Architecture range across five models': (f'{arch2:.2f}', 'no'),
        'Training-target range against G₂': (f'{sup2:.2f}', 'no'),
        'ResNet-18 versus ResNet-34': (f'{r18_34:.2f}', 'no'),
        'C − B (soft supervision) against G₂': (f'{cb2:.2f}', 'no'),
    }
    for row in tb[6].rows[1:]:
        lab = row.cells[0].text.strip()
        if lab in T6:
            setcell(row.cells[1], T6[lab][0])
            setcell(row.cells[2], T6[lab][1])

    # ---------- Table 7：测试集矩阵 ----------
    A7 = ['tgtA', 'tgtB', 'tgtC', 'tgtD', 'tgtE', 'archR18', 'archCNX', 'archPVT', 'archPU']
    for row, arm in zip(tb[7].rows[1:], A7):
        for j, k in enumerate([1, 2, 3, 4], start=1):
            e = t1t[arm][str(k)]
            setcell(row.cells[j], f"{e['dice_05']:.4f} / {e['dice_opt']:.4f}")

    # ---------- Table 8：跨数据集推广 ----------
    IDX = {'LIDC-IDRI': 'LIDC', 'QUBIQ brain-growth': 'QUBIQ-brain', 'QUBIQ kidney': 'QUBIQ-kidney'}
    rows = tb[8].rows
    cur = None
    for r in rows[1:]:
        lab = r.cells[0].text.strip()
        if lab in IDX:
            cur = IDX[lab]
            setcell(r.cells[2], f"{DIS[cur]['area_mean_ratio']:.2f}")
        assert cur, '表 8 行顺序异常'
        mode = 'fixed0.5' if r.cells[3].text.strip().startswith('fixed') else 'opt_thr'
        e = G[cur][mode]
        setcell(r.cells[4], f"{e['protocol']:.2f}")
        setcell(r.cells[5], f"{e['supervision']:.2f}")
        setcell(r.cells[6], f"{e['ratio']:.2f}×")

    # ---------- 插图替换 ----------
    NEW = {2: 'fig2_threshold_sensitivity.png', 3: 'fig4_target_matrix.png',
           4: 'fig5_forest.png', 5: 'fig3_effect_sizes.png',
           6: 'fig6_calibration.png', 9: 'fig9_generalisation.png'}
    imgs = [p for p in doc.paragraphs if 'graphicData' in p._element.xml]
    assert len(imgs) == 9, f'图数量异常: {len(imgs)}'
    for n, par in enumerate(imgs, 1):
        if n not in NEW:
            continue
        sh = par._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline')
        cx = cy = None
        for s in doc.inline_shapes:
            if s._inline in sh:
                cx, cy = s.width, s.height
                break
        for r in list(par.runs):
            r._element.getparent().remove(r._element)
        run = par.add_run()
        if cx:
            run.add_picture(os.path.join(FIG, NEW[n]), width=Emu(cx))
        else:
            run.add_picture(os.path.join(FIG, NEW[n]))

    doc.save(SRC)

    d2 = docx.Document(SRC)
    print(f'表格与插图已更新；图 {len(d2.inline_shapes)} 张，表 {len(d2.tables)} 个')
    print('Table 2 首行:', ' | '.join(c.text for c in d2.tables[2].rows[1].cells))
    print('Table 8 首行:', ' | '.join(c.text for c in d2.tables[8].rows[1].cells))
    print(SRC)


if __name__ == '__main__':
    main()
