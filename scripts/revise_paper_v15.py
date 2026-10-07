# -*- coding: utf-8 -*-
"""revise_paper_v15.py — v14 -> v15

依据: 2026-09-04 用户拍板的两项决策, 加本轮新查出的一处渲染缺陷。

修改清单:
  1. [渲染缺陷] 删除首段残留的中文镜像抬头。该段是 Title 样式的正文段,
     会作为论文渲染后的第一行印出来 (已在 v14_preview.pdf 第 1 页确认)。
  2. [真实性] §3.1: 53.7%/0.7% 的"检出层面 vs 勾画层面"断言无脚本支撑,
     且投票图在结构上无法区分两者 (四人各自标注但无公共像素时 maxV<4,
     与"只有三人标注"同形)。改为用论文自身实例定义可复算的比例
     4958/10533 = 47.1% (源 outputs/analysis/dataset_stats.json 的
     gt_levels 计数, 与同段已有的 4579/8471 与 996/2062 算术自洽),
     并明写该区分不可从投票图恢复。
  3. [篇幅] 附录 A、附录 B (含 Table S1) 移出正文另存补充材料,
     改称 Supplementary Note S1 / S2。§4.7 保留正文, 全文不重编号,
     图号不变 (Fig. 1-9 与仓库 fig1-fig9 文件名保持一一对应)。
  4. 交叉引用同步: §3.4 的 "Appendix A" -> "Supplementary Note S1";
     §5.1 增补一句指向 Table S1 (原附录 B 在正文中无任何引用)。

用法: python scripts/revise_paper_v15.py
输出: outputs/paper/论文_协议效应_v15_CIBM.docx
      outputs/paper/论文_协议效应_v15_补充材料.docx
"""
import copy
import os

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAPER = os.path.join(ROOT, 'outputs', 'paper')
SRC = os.path.join(PAPER, '论文_协议效应_v14_CIBM.docx')
DST = os.path.join(PAPER, '论文_协议效应_v15_CIBM.docx')
SUP = os.path.join(PAPER, '论文_协议效应_v15_补充材料.docx')

TITLE = ('Protocol outweighs method: patch sampling, evaluation scope, consensus level '
         'and threshold dominate architecture and supervision-target choice in lung '
         'nodule segmentation on LIDC-IDRI')

# ---- 1. 机制段落: 精确到 run 级替换, 保留 G/V 的斜体 run ----
OLD_MECH = (
    'We checked the mechanism behind this. Over the nodule instances that at least two '
    'readers marked, 53.7 per cent were marked by only two or three of the four readers '
    '— detection-level disagreement, in which the lesion is simply absent from the '
    'other readers’ mark-ups — and only 0.7 per cent were marked by all four '
    'with no pixel in common. The empty high-consensus references are therefore dominated '
    'by detection-level disagreement, not by contour-level disagreement among readers who '
    'all marked the lesion. ')

NEW_MECH = (
    'Pooling the two sets, 4958 of the 10533 instances — 47.1 per cent — have no '
    'pixel on which all four readers agree. What the vote map cannot supply is the reason. '
    'A lesion that only two or three readers marked at all, and a lesion that all four '
    'marked but whose contours share no single common pixel, both reduce to a maximum vote '
    'count below four; once the per-reader masks have been collapsed into a count, the two '
    'are indistinguishable. Separating them would require the individual readers’ '
    'contours aligned to the reconstructed slice index, which the vote-map representation '
    'does not retain. We therefore report the proportion as a property of the reference '
    'standard and draw no conclusion about which kind of disagreement produces it; the '
    'distinction is not recoverable from the representation in which multi-rater ground '
    'truth is conventionally stored. ')

OLD_APX_A = 'reported in Appendix A as a negative result'
NEW_APX_A = 'reported in Supplementary Note S1 as a negative result'

S1_POINTER = (' A same-protocol comparison of this study’s arms with the closest '
              'published figures, together with the four protocol quantities behind each, '
              'is given as Table S1 in Supplementary Note S2.')

APX_A_HEAD = 'Appendix A. The earlier six-term loss, reported as a negative result'
APX_B_HEAD = 'Appendix B. Same-protocol baseline comparison'
NEW_A_HEAD = 'Supplementary Note S1. The earlier six-term loss, reported as a negative result'
NEW_B_HEAD = 'Supplementary Note S2. Same-protocol baseline comparison'
BLOCK_LAST = 'Third, the protocol columns are not always fillable from the published text.'


def run_replace(par, old, new, label):
    """在段落的某个 run 内做纯字符串替换, 不触碰其它 run 的格式。"""
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    raise AssertionError(f'{label}: 未在任何 run 中找到目标文本')


def find_par(doc, pred, label):
    for p in doc.paragraphs:
        if pred(p.text):
            return p
    raise AssertionError(f'{label}: 未找到目标段落')


def block_slice(doc):
    """返回附录块在 body children 中的 [start, end] 闭区间下标（含中间的表格）。"""
    body = doc.element.body
    kids = list(body)
    pa = find_par(doc, lambda t: t.strip() == APX_A_HEAD, 'Appendix A 标题')
    pz = find_par(doc, lambda t: t.strip().startswith(BLOCK_LAST), '附录 B 末段')
    i, j = kids.index(pa._element), kids.index(pz._element)
    assert i < j, '附录块首尾顺序异常'
    return body, kids, i, j


def build_main():
    doc = docx.Document(SRC)
    ps = doc.paragraphs

    # 1. 删除残留的中文镜像抬头
    assert '镜像' in ps[0].text and '非权威源' in ps[0].text, f'首段不是镜像抬头: {ps[0].text!r}'
    ps[0]._element.getparent().remove(ps[0]._element)
    # 抬头之后残留的空段落一并清掉, 使文档以标题起首
    while doc.paragraphs and not doc.paragraphs[0].text.strip():
        el = doc.paragraphs[0]._element
        el.getparent().remove(el)

    # 2. §3.1 机制段
    p_mech = find_par(doc, lambda t: '53.7 per cent' in t, '§3.1 机制段')
    run_replace(p_mech, OLD_MECH, NEW_MECH, '§3.1 机制句')

    # 3. §3.4 Appendix A 引用
    p_loss = find_par(doc, lambda t: OLD_APX_A in t, '§3.4 Appendix A 引用')
    run_replace(p_loss, OLD_APX_A, NEW_APX_A, 'Appendix A 引用')

    # 4. §5.1 增补 Table S1 指针
    p_51 = find_par(
        doc,
        lambda t: t.strip().startswith('It follows that a difference of one to three Dice points'),
        '§5.1 指针锚点')
    p_51.runs[-1].text = p_51.runs[-1].text + S1_POINTER

    # 5. 移除附录块
    body, kids, i, j = block_slice(doc)
    for el in kids[i:j + 1]:
        body.remove(el)

    doc.save(DST)
    return doc


def build_supplement():
    doc = docx.Document(SRC)
    body, kids, i, j = block_slice(doc)
    keep = set(id(e) for e in kids[i:j + 1])
    tag_sect = '}sectPr'
    for el in kids:
        if id(el) in keep or el.tag.endswith(tag_sect):
            continue
        body.remove(el)

    # 重命名两个标题 + 首句"appendix"->"note"
    ha = find_par(doc, lambda t: t.strip() == APX_A_HEAD, '补充材料 A 标题')
    run_replace(ha, APX_A_HEAD, NEW_A_HEAD, '补充材料 A 标题')
    hb = find_par(doc, lambda t: t.strip() == APX_B_HEAD, '补充材料 B 标题')
    run_replace(hb, APX_B_HEAD, NEW_B_HEAD, '补充材料 B 标题')
    pp = find_par(doc, lambda t: t.strip().startswith('The purpose of this appendix is narrow'),
                  '附录 B 首句')
    run_replace(pp, 'The purpose of this appendix is narrow',
                'The purpose of this note is narrow', '附录 B 首句')

    # 顶部插入标题页
    first = doc.paragraphs[0]._element
    for text, style in ((TITLE, None), ('Supplementary material', 'Title')):
        p = doc.add_paragraph(text)
        if style:
            p.style = doc.styles[style]
        first.addprevious(p._element)

    doc.save(SUP)
    return doc


def main():
    build_main()
    build_supplement()

    # ---- 回读自检 ----
    m = docx.Document(DST)
    mt = '\n'.join(p.text for p in m.paragraphs)
    for bad in ('53.7', '0.7 per cent', 'detection-level disagreement',
                'Appendix A', 'Appendix B', '镜像', '非权威源'):
        assert bad not in mt, f'v15 正文仍含: {bad}'
    for good in ('4958 of the 10533 instances', '47.1 per cent',
                 'Supplementary Note S1', 'Table S1 in Supplementary Note S2'):
        assert good in mt, f'v15 正文缺少: {good}'
    assert m.paragraphs[0].text.strip().startswith('Protocol outweighs method'), \
        f'v15 首段应为标题, 实为: {m.paragraphs[0].text[:60]!r}'

    s = docx.Document(SUP)
    st = '\n'.join(p.text for p in s.paragraphs)
    for good in (NEW_A_HEAD, NEW_B_HEAD, 'Table S1.', 'The purpose of this note is narrow'):
        assert good in st, f'补充材料缺少: {good}'
    assert 'Appendix' not in st, '补充材料仍含 Appendix 字样'

    print('=== v15 ===')
    print(f'  段落 {len(m.paragraphs)}  表格 {len(m.tables)}  图 {len(m.inline_shapes)}')
    print(f'  正文词数 {sum(len(p.text.split()) for p in m.paragraphs)}')
    print(f'  PENDING {mt.count("PENDING")}')
    print('=== 补充材料 ===')
    print(f'  段落 {len(s.paragraphs)}  表格 {len(s.tables)}  图 {len(s.inline_shapes)}')
    print(f'  词数 {sum(len(p.text.split()) for p in s.paragraphs)}')
    print('回读自检全部通过')
    print(DST)
    print(SUP)


if __name__ == '__main__':
    main()
