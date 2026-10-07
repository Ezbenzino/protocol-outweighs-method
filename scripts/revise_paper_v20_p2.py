# -*- coding: utf-8 -*-
"""revise_paper_v20_p2.py —— v20 第二轮审计修复（不升版本号，权威源仍是 v20 docx）。

本轮修三处**可客观判定为错**的项，不动任何论点：

 1. **Table 9 第 4–17 行没有套用表格格式**。前 3 行（表头 + LIDC 两行）是
    8.5 pt、数值列右对齐、带单元格内边距；QUBIQ 扩写时新增的 14 行**全部没有**，
    渲染出来字号明显偏大、整列左对齐（见 preview PDF 第 29 页）。
    九张表里只有这一张有这个问题。

 2. **§4.5 的 "14.1×" 应为 "14.0×"**。位置不变臂的协议/监督比值
    = 43.837686 / 3.1240 = **14.03**。14.1 是先把分子分母各自四舍五入
    （43.84 / 3.12 = 14.05）再取一位小数得到的，属二次舍入。
    同段另外三个比值都用未舍入值算，口径不一致：
      22.510908 / 3.1240   = 7.2058  -> 7.2×  ✓
      22.510908 / 1.873683 = 12.014  -> 12.0× ✓
      43.837686 / 1.873683 = 23.396  -> 23.4× ✓

 3. **Fig. 5（森林图）全文从未被引用**。Elsevier 要求每张图表都在正文中被引用。
    在它所属的 §4.3 末尾补一处引用。
    ⚠️ 补完之后首次引用顺序变成 Fig. 3 → Fig. 5 → Fig. 4，与图号不一致；
    彻底解决需要把 Fig. 4/Fig. 5 对调编号并交换两个图块的位置，
    那会改变图的身份（答辩 PPT / 中文手册里都按现编号引用），故留给作者决定，
    本脚本不做。Table 3/4 也存在同类顺序问题（Table 4 在 §3 先被引用）。

用法：
    python scripts/revise_paper_v20_p2.py [in.docx] [out.docx]
"""
import re
import shutil
import sys

from docx import Document
from docx.oxml.ns import qn

SRC = sys.argv[1] if len(sys.argv) > 1 else 'outputs/paper/论文_协议效应_v20_CIBM.docx'
DST = sys.argv[2] if len(sys.argv) > 2 else SRC

FIG5_CITE = (' Fig. 5 shows all three contrasts against each evaluation reference, '
             'with the run-to-run noise floor marked.')


def set_cell_format(cell, align, half_pt=17):
    """把 Table 9 前三行的格式套到一个单元格上：内边距 + 行距 + 对齐 + 字号。"""
    tcPr = cell._tc.get_or_add_tcPr()
    if tcPr.find(qn('w:tcMar')) is None:
        mar = tcPr.makeelement(qn('w:tcMar'), {})
        for tag, w in (('w:top', '60'), ('w:left', '90'),
                       ('w:bottom', '60'), ('w:right', '90')):
            e = mar.makeelement(qn(tag), {qn('w:type'): 'dxa', qn('w:w'): w})
            mar.append(e)
        tcPr.append(mar)
    for p in cell.paragraphs:
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn('w:spacing')) is None:
            sp = pPr.makeelement(qn('w:spacing'),
                                 {qn('w:after'): '0', qn('w:line'): '240'})
            pPr.insert(0, sp)
        jc = pPr.find(qn('w:jc'))
        if jc is None:
            jc = pPr.makeelement(qn('w:jc'), {})
            pPr.append(jc)
        jc.set(qn('w:val'), align)
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            for tag in ('w:sz', 'w:szCs'):
                e = rPr.find(qn(tag))
                if e is None:
                    e = rPr.makeelement(qn(tag), {})
                    rPr.append(e)
                e.set(qn('w:val'), str(half_pt))


def main():
    if DST != SRC:
        shutil.copyfile(SRC, DST)
    d = Document(DST)

    # ---- 1. Table 9 行格式 ----
    t9 = d.tables[8]
    assert len(t9.rows) == 17 and t9.rows[0].cells[0].text.strip() == 'Dataset', \
        'tables[8] 不是 Table 9，先确认表序'
    fixed = 0
    for ri in range(3, len(t9.rows)):
        for ci, cell in enumerate(t9.rows[ri].cells):
            set_cell_format(cell, 'left' if ci == 0 else 'right')
            fixed += 1
    print(f'  Table 9：补齐 {fixed} 个单元格的格式（第 4–17 行）')

    # ---- 2. 14.1× -> 14.0× ----
    hit = 0
    for p in d.paragraphs:
        if '14.1×' in ''.join(r.text for r in p.runs):
            for r in p.runs:
                if '14.1×' in r.text:
                    r.text = r.text.replace('14.1×', '14.0×'); hit += 1; break
            else:
                full = ''.join(r.text for r in p.runs)
                p.runs[0].text = full.replace('14.1×', '14.0×')
                for r in p.runs[1:]:
                    r.text = ''
                hit += 1
    assert hit == 1, f'14.1× 命中 {hit} 次，预期 1 次'
    print('  §4.5：14.1× -> 14.0×')

    # ---- 3. 补 Fig. 5 引用 ----
    p136 = d.paragraphs[136]
    assert p136.text.startswith('Consequently, the apparent advantage'), \
        f'第 136 段不是 §4.3 收尾段，实为 {p136.text[:60]!r}'
    p136.runs[-1].text = p136.runs[-1].text + FIG5_CITE
    print('  §4.3：补 Fig. 5 引用')

    d.save(DST)

    # ================= 回读自检 =================
    d2 = Document(DST)
    txt = '\n'.join(p.text for p in d2.paragraphs)
    assert '14.1×' not in txt and '14.0×' in txt, '比值未改成功'
    assert 'Fig. 5 shows all three contrasts' in txt, 'Fig. 5 引用未写入'
    t9 = d2.tables[8]
    bad = 0
    for ri in range(len(t9.rows)):
        for c in t9.rows[ri].cells:
            for p in c.paragraphs:
                if p.alignment is None:
                    bad += 1
                for r in p.runs:
                    if r.font.size is None or abs(r.font.size.pt - 8.5) > 1e-6:
                        bad += 1
    assert bad == 0, f'Table 9 仍有 {bad} 处未套格式'
    # 每张图表都被引用
    capre = re.compile(r'^\s*(Table|Fig\.)\s*(\d+)\s*[.—–-]')
    paras = [p.text for p in d2.paragraphs]
    caps = {}
    for i, s in enumerate(paras):
        m = capre.match(s.strip())
        if m:
            caps[(m.group(1), int(m.group(2)))] = i
    EX = {i for i, s in enumerate(paras) if 'their Table' in s or 'their Section 3.6.2' in s}
    cited = set()
    for i, s in enumerate(paras):
        if i in EX:
            continue
        m = capre.match(s.strip())
        own = (m.group(1), int(m.group(2))) if m else None
        for mm in re.finditer(r'\b(Table|Fig\.)\s*(\d+)', s):
            k = (mm.group(1), int(mm.group(2)))
            if k != own:
                cited.add(k)
    missing = sorted(set(caps) - cited)
    print(f'  自检通过：{len(caps)} 个图表标题，未被引用的 = {missing or "无"}')
    print('->', DST)


if __name__ == '__main__':
    main()
