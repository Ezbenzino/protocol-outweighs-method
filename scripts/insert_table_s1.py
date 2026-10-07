"""Insert Table S1 into v10 docx (restore from v9) + upgrade Zhi threshold claim.
python-docx edit."""
import docx, copy, sys
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'outputs/paper/论文_协议效应_v10.docx'

doc = docx.Document(SRC)

# ---- locate the placeholder paragraph ----
target = None
idx = -1
for i, p in enumerate(doc.paragraphs):
    if 'Retained from the previous version' in p.text:
        target = p; idx = i; break
if target is None:
    sys.exit('ERROR: Table S1 placeholder not found')

# ---- intro sentence (replaces placeholder) ----
target.text = ('DSC is reported on a scale of 0\u20131 (this work) or 0\u2013100 (Zhi et al. [REF 39]). '
               'Protocol, not method, is the dominant source of between-study variation.')

# ---- build the table ----
rows = [
    ['Method / study', 'Reported DSC',
     'Evaluation protocol (fusion / slice / localisation / threshold)', 'Comparability'],
    ['This work, arm B (majority-vote training, ResNet-34)', '0.8675 (V \u2265 2)',
     'Majority of 4 readers / 2-D single slice / oracle-centred 128 px patch / calibrated threshold',
     'Same-protocol baseline'],
    ['This work, arm D (soft-vote training)', '0.8686 (V \u2265 2)',
     'Same as above', 'Same protocol'],
    ['This work, five architectures (four pretrained)', '0.8664\u20130.8692 (V \u2265 2)',
     'Same as above', 'Same protocol'],
    ['Zhi et al. [REF 39], five best 2-D LIDC models', '84.23\u201383.92 (DSC \u00d7 100)',
     'Mean of 4 reader masks / largest-cross-section centre-crop 64 px / threshold fixed at 0.5 in the released code / no repeated runs',
     'Protocol differs \u2014 reference only'],
    ['Zhi et al. [REF 39], loss effect (Dice \u2192 BCE \u2192 Focal)', '82.55 \u2192 79.78 \u2192 65.35',
     'Same as above', 'Protocol differs \u2014 reference only'],
]

# insert a table after the target paragraph
tbl = doc.add_table(rows=len(rows), cols=4)
# add explicit borders (no 'Table Grid' style in this doc)
tblPr = tbl._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '000000')
    borders.append(el)
tblPr.append(borders)
for ri, rowdata in enumerate(rows):
    for ci, val in enumerate(rowdata):
        tbl.rows[ri].cells[ci].text = val
# make header bold
for ci in range(4):
    for run in tbl.rows[0].cells[ci].paragraphs[0].runs:
        run.font.bold = True

# move the new table to right after the target paragraph
target._p.addnext(tbl._tbl)

doc.save(SRC)
print(f'OK: placeholder paragraph P{idx} replaced, Table S1 inserted ({len(rows)} rows x 4 cols), saved {SRC}')
