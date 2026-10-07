# -*- coding: utf-8 -*-
"""统一论文图/表编号：按正文出现顺序重编号，修复 Table 1/4 重复与 Fig 3/4/5 颠倒。
外部文献引用（their Table/Fig）不动。每项替换校验唯一出现。
"""
import docx, re, sys

PATH = r'outputs/paper/论文初稿_协议效应_v9.docx'
d = docx.Document(PATH)

# (源子串, 新子串) —— 全部基于当前 docx 段落文本确认的唯一子串
REPL = [
    # ---- 表标题（编号 1->2,2->3,3->5,4->6,5->7,6->8,7->9；Cohort=1、Effects=4 不变）
    ("Table 1. Cross-evaluation matrix", "Table 2. Cross-evaluation matrix"),
    ("Table 2. Case-level paired comparisons", "Table 3. Case-level paired comparisons"),
    ("Table 3. Architecture control", "Table 5. Architecture control"),
    ("Table 4. Held-out test set", "Table 6. Held-out test set"),
    ("Table 5. Dice against G\u2082 as the evaluation window is displaced",
     "Table 7. Dice against G\u2082 as the evaluation window is displaced"),
    ("Table 6. Dice against G\u2082 as the evaluation scope widens",
     "Table 8. Dice against G\u2082 as the evaluation scope widens"),
    ("Table 7. Protocol (evaluation-target)", "Table 9. Protocol (evaluation-target)"),
    # ---- 正文表引用
    ("under the same conditions (Table 1)", "under the same conditions (Table 2)"),
    ("is given in Table 3", "is given in Table 5"),
    ("Every conclusion replicated (Table 4)", "Every conclusion replicated (Table 6)"),
    ("Table 5 reports the same measurement", "Table 7 reports the same measurement"),
    ("Two entries of Table 6 require", "Two entries of Table 8 require"),
    ("column of Table 6 was qualified", "column of Table 8 was qualified"),
    # ---- 图标题（Fig 4->3, Fig 5->4, Fig 3->5；其余不变）
    ("Fig. 4 \u2014 Training-target \u00d7 evaluation-target cross-evaluation matrix",
     "Fig. 3 \u2014 Training-target \u00d7 evaluation-target cross-evaluation matrix"),
    ("Fig. 5 \u2014 Forest plot", "Fig. 4 \u2014 Forest plot"),
    ("Fig. 3 \u2014 Effect sizes of protocol", "Fig. 5 \u2014 Effect sizes of protocol"),
    # ---- 正文图引用
    ("near-diagonal structure (Fig. 4)", "near-diagonal structure (Fig. 3)"),
    ("Fig. 3 places all measured effects", "Fig. 5 places all measured effects"),
    ("(Fig. 3, shaded band)", "(Fig. 5, shaded band)"),
]

def all_text(doc):
    """遍历段落 + 表格 cell，返回 (容器, 文本) 列表，容器可被替换。"""
    out = []
    for p in doc.paragraphs:
        out.append(('para', p))
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    out.append(('cell', p))
    return out

def replace_in_runs(p, old, new):
    """在段落 run 层面替换（避免整段重建丢格式）。"""
    # 简单情形：直接替换 paragraph.text 需要逐 run。这里用逐 run 拼接策略：
    # 若某 run 含 old，直接替换该 run；否则跨 run 场景少见，此处整段重建前先尝试 run 级。
    found = False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            found = True
    if found:
        return 1
    # 跨 run 场景：整段文本替换（保留首 run 格式，其余清空）
    full = p.text
    if old in full:
        new_full = full.replace(old, new)
        if p.runs:
            p.runs[0].text = new_full
            for r in p.runs[1:]:
                r.text = ''
        else:
            p.add_run(new_full)
        return 1
    return 0

stats = []
containers = all_text(d)
for old, new in REPL:
    cnt = 0
    for kind, p in containers:
        cnt += replace_in_runs(p, old, new)
    stats.append((old, cnt))
    if cnt != 1:
        print(f'⚠️  {"MISS" if cnt==0 else "MULTI"}({cnt}): {old[:70]}')

d.save(PATH)
print('保存完成。校验：')
allok = True
for old, cnt in stats:
    mark = 'OK ' if cnt == 1 else 'FAIL'
    if cnt != 1: allok = False
    print(f'  {mark} [{cnt}] {old[:75]}')
print('全部唯一替换通过' if allok else '存在未通过项，请检查')
