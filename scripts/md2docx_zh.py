# -*- coding: utf-8 -*-
"""md -> docx 转换器：支持 # 标题、**加粗**、表格、列表"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_ea(run, font='宋体'):
    run.font.name = 'Times New Roman'
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)

def add_runs(p, text, size=10.5, bold=False, italic=False, color=None, font='宋体'):
    # 解析 **bold** 与 *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for tok in tokens:
        if not tok:
            continue
        b, it = bold, italic
        t = tok
        if tok.startswith('**') and tok.endswith('**'):
            t = tok[2:-2]; b = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            t = tok[1:-1]; it = True
        r = p.add_run(t)
        r.font.size = Pt(size)
        r.bold = b
        r.italic = it
        if color: r.font.color.rgb = RGBColor(*color)
        set_ea(r, font)

def convert(md_path, docx_path, title_style=True):
    doc = Document()
    # Normal 样式
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    lines = open(md_path, encoding='utf-8').read().split('\n')
    i = 0
    in_table = False
    table_rows = []
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1; continue
        # 表格行
        if stripped.startswith('|') and stripped.endswith('|'):
            # 分隔行
            if re.match(r'^\|[\s\-:|]+\|$', stripped) and '-' in stripped:
                i += 1; continue
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            table_rows.append(cells)
            # 检查下一行是否还是表格
            j = i + 1
            nxt = lines[j].rstrip() if j < len(lines) else ''
            if not (nxt.startswith('|') and nxt.endswith('|')):
                # 渲染表格
                ncol = max(len(r) for r in table_rows)
                tb = doc.add_table(rows=len(table_rows), cols=ncol)
                tb.style = 'Table Grid'
                for ri, row in enumerate(table_rows):
                    for ci in range(ncol):
                        cell_text = row[ci] if ci < len(row) else ''
                        cell = tb.rows[ri].cells[ci]
                        cell.text = ''
                        p = cell.paragraphs[0]
                        add_runs(p, cell_text, size=9.5, bold=(ri == 0))
                doc.add_paragraph()
                table_rows = []
            i += 1; continue
        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                p = doc.add_heading(level=0)
                add_runs(p, text, size=16, bold=True, font='黑体')
            elif level == 2:
                p = doc.add_heading(level=1)
                add_runs(p, text, size=14, bold=True, font='黑体')
            elif level == 3:
                p = doc.add_heading(level=2)
                add_runs(p, text, size=12, bold=True, font='黑体')
            else:
                p = doc.add_heading(level=3)
                add_runs(p, text, size=11, bold=True, font='黑体')
            i += 1; continue
        # 列表
        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, stripped[2:])
            i += 1; continue
        if re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, stripped)
            i += 1; continue
        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    print('已生成:', docx_path)

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
