# -*- coding: utf-8 -*-
"""revise_paper_v20.py —— v19 -> v20：冻结后审核发现的缺陷修复。

本轮**不新增任何论点**，只修被审核出来的错：

 1. Fig. 3 图片：v16 那轮换图时没核对数值，图上一直是**已废弃的实例级**
    16.42 / 10.02 / 2.99 / 1.22 / 1.15，与正文、Table 5 和它自己的图注全部矛盾。
    换成 scripts/make_fig3_ranked.py 现算的病例级测试集版本。
 2. Fig. 9 图片：图上 "permutation p = 0.048" 是**写死的字面量**，且是 7 点的 p；
    8 点（含 LIDC）的精确置换 p 是 0.011，图注与 §4.14 都写 0.011。
    另有两处子图标签严重重叠。换成 scripts/make_fig9_generalisation.py 重画的版本。
 3. Table 9 **没有标题**（只有 "Table 9." 四个字，且未加粗，与其余八张表不一致）。
 4. §4.14 读者数列表 "(3, 3, 3, 6, 6 and 7)" 少一个 3——七个 QUBIQ 任务的读者数是
    3, 3, 3, 3, 6, 6, 7（Table 9 可直接核对），§5.5 的八点版本写的是对的。
 5. Table 1 / Table 6 / Fig. 6 / Fig. 8 **全文从未被引用**。补四处引用，
    插入位置都不改变"编号顺序 = 首次引用顺序"。
    （Fig. 5 未被引用一并查出，但补它会打乱图号顺序，留给作者决定，本轮不动。）
 6. 删 7 个孤儿图片部件（docx 里 16 个 media，只有 9 个在用），约省 2.2 MB。

图片替换时按新图的宽高比重算 extent，避免被拉伸。
"""
import io
import os
import re
import shutil
import struct
import zipfile

from docx import Document
from docx.shared import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'outputs', 'paper', '论文_协议效应_v19_CIBM.docx')
DST = os.path.join(ROOT, 'outputs', 'paper', '论文_协议效应_v20_CIBM.docx')
FIGS = os.path.join(ROOT, 'outputs', 'figures')

TABLE9_CAPTION = (
    ' External replication on seven QUBIQ 2021 tasks, with LIDC-IDRI as the reference '
    'point. The agreement index is Σ|V ≥ N| / Σ|V ≥ 1|, computed on the '
    'slices that enter training and evaluation (Section 3.8). The protocol effect is the '
    'largest Dice gap across the supervision arms between the V ≥ 1 and V ≥ N '
    'references; the supervision effect is the largest range across those arms at any single '
    'reference. Every task is reported twice, at a fixed threshold of 0.5 and at each arm’s '
    'calibrated optimum; all values are five-fold cross-validation estimates with the case as '
    'the unit of analysis. The ratio column is unstable where the supervision effect approaches '
    'zero and should be read alongside the absolute effects (Section 4.14).')

# (段落索引, 旧片段, 新片段)   旧片段为 None 表示"追加到段尾"
EDITS = [
    (57, None, ' Table 1 summarises the cohort composition and the nodule size distribution.'),
    (162, '0.96 Dice points at V ≥ 2, between',
          '0.96 Dice points at V ≥ 2 (Table 6), between'),
    (166, None, ' Fig. 6 plots the mean predicted area against the threshold and the Brier '
                'score of each arm.'),
    (207, None, ' Fig. 8 shows the whole ladder for four arms differing only in their '
                'training-side treatment of position and background.'),
    (232, '(3, 3, 3, 6, 6 and 7)', '(3, 3, 3, 3, 6, 6 and 7)'),
]


def para_replace(p, old, new):
    """在段落内做一次精确替换，尽量保留 run 级格式。"""
    full = ''.join(r.text for r in p.runs)
    if old not in full:
        raise AssertionError(f'未找到待替换片段: {old!r}')
    # 优先在单个 run 内替换
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return
    # 跨 run：把整段压回第一个 run
    p.runs[0].text = full.replace(old, new, 1)
    for r in p.runs[1:]:
        r.text = ''


def para_append(p, text):
    p.runs[-1].text = p.runs[-1].text + text


def png_size(b):
    assert b[:8] == b'\x89PNG\r\n\x1a\n'
    return struct.unpack('>II', b[16:24])


def main():
    shutil.copyfile(SRC, DST)
    d = Document(DST)
    paras = d.paragraphs

    for idx, old, new in EDITS:
        if old is None:
            para_append(paras[idx], new)
        else:
            para_replace(paras[idx], old, new)

    # ---- Table 9 标题：加粗标签 + 补正文 ----
    p9 = paras[235]
    assert p9.text.strip() == 'Table 9.', f'第 235 段不是空标题，实为 {p9.text!r}'
    r0 = p9.runs[0]
    r0.bold = True
    r_new = p9.add_run(TABLE9_CAPTION)
    r_new.bold = False

    # ---- 图片按新宽高比重设 extent ----
    newpix = {}
    for rid, fname in (('rId20', 'fig3_effect_sizes.png'),
                       ('rId23', 'fig9_generalisation.png')):
        with open(os.path.join(FIGS, fname), 'rb') as f:
            newpix[rid] = png_size(f.read())
    for shp in d.inline_shapes:
        rid = shp._inline.graphic.graphicData.pic.blipFill.blip.embed
        if rid in newpix:
            w, h = newpix[rid]
            cx = shp.width
            shp.height = Emu(int(round(int(cx) * h / w)))
            print(f'  {rid}: extent -> cx={int(cx)} cy={int(shp.height)}  ({w}x{h})')
    d.save(DST)

    # ---- 换图片字节 + 删孤儿部件 ----
    zin = zipfile.ZipFile(DST)
    doc = zin.read('word/document.xml').decode('utf-8')
    rels = zin.read('word/_rels/document.xml.rels').decode('utf-8')
    relmap = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    used_rids = set(re.findall(r'r:embed="([^"]+)"', doc))
    used_media = {'word/' + relmap[r] for r in used_rids if r in relmap}
    orphans = [n for n in zin.namelist()
               if n.startswith('word/media/') and n not in used_media]
    swap = {'word/' + relmap['rId20']: os.path.join(FIGS, 'fig3_effect_sizes.png'),
            'word/' + relmap['rId23']: os.path.join(FIGS, 'fig9_generalisation.png')}

    orphan_ids = [i for i, t in relmap.items()
                  if t.startswith('media/') and 'word/' + t in orphans]
    for i in orphan_ids:
        rels = re.sub(r'<Relationship Id="%s"[^>]*/>' % re.escape(i), '', rels)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            n = item.filename
            if n in orphans:
                continue
            if n == 'word/_rels/document.xml.rels':
                zout.writestr(item, rels.encode('utf-8'))
            elif n in swap:
                with open(swap[n], 'rb') as f:
                    zout.writestr(item, f.read())
            else:
                zout.writestr(item, zin.read(n))
    zin.close()
    with open(DST, 'wb') as f:
        f.write(buf.getvalue())
    print(f'  删除孤儿图片 {len(orphans)} 个，移除 rels 条目 {len(orphan_ids)} 条')

    # ================= 回读自检 =================
    d2 = Document(DST)
    txt = '\n'.join(p.text for p in d2.paragraphs)
    tab = '\n'.join(c.text for t in d2.tables for r in t.rows for c in r.cells)
    allt = txt + '\n' + tab

    checks_present = [
        'Table 1 summarises the cohort composition',
        '0.96 Dice points at V ≥ 2 (Table 6)',
        'Fig. 6 plots the mean predicted area',
        'Fig. 8 shows the whole ladder',
        '(3, 3, 3, 3, 6, 6 and 7)',
        'External replication on seven QUBIQ 2021 tasks',
    ]
    checks_absent = [
        '(3, 3, 3, 6, 6 and 7)',
        'Agreement inde×', 'p = 0.007', 'falls monotonically',
    ]
    for s in checks_present:
        assert s in allt, f'自检失败：应存在但未找到 {s!r}'
    for s in checks_absent:
        assert s not in allt, f'自检失败：应删除但仍存在 {s!r}'

    # Table 9 标题不再为空
    for i, p in enumerate(d2.paragraphs):
        if p.text.strip().startswith('Table 9.'):
            assert len(p.text.strip()) > 60, 'Table 9 标题仍然是空的'
            break

    # 图片字节确实换了
    z = zipfile.ZipFile(DST)
    for rid, fname in (('rId20', 'fig3_effect_sizes.png'),
                       ('rId23', 'fig9_generalisation.png')):
        tgt = 'word/' + relmap[rid]
        assert z.read(tgt) == open(os.path.join(FIGS, fname), 'rb').read(), \
            f'{rid} 图片未替换成功'
    media = [n for n in z.namelist() if n.startswith('word/media/')]
    embeds = set(re.findall(r'r:embed="([^"]+)"',
                            z.read('word/document.xml').decode('utf-8')))
    print(f'  自检通过：media 部件 {len(media)} 个，文档内嵌 {len(embeds)} 个')
    print('->', DST)


if __name__ == '__main__':
    main()
