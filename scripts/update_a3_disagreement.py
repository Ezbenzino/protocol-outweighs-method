import docx
PATH = r'outputs/paper/论文初稿_协议效应_v9.docx'
doc = docx.Document(PATH)

def rep(p, old, new, label):
    if old in p.text:
        t = p.text
        nt = t.replace(old, new)
        if p.runs:
            p.runs[0].text = nt
            for r in p.runs[1:]:
                r.text = ''
        print(f'{label}: 已替换')
        return True
    print(f'{label}: 未找到!')
    return False

for p in doc.paragraphs:
    # [237] 分歧指数定义 + 65%
    rep(p,
        'Define a disagreement index as |V \u2265 N| / |V \u2265 1|, the fraction of the union mask on which all raters agree; smaller values mean greater inter-rater disagreement. Across LIDC, QUBIQ brain-growth and QUBIQ kidney the index is 0.18, 0.50 and 0.87, and the protocol-to-supervision ratio at optimal thresholds is 2.69\u00d7, 1.44\u00d7 and 0.78\u00d7 respectively. The three points are strictly monotone: the greater the rater disagreement, the more the protocol choice dominates the method choice; as agreement approaches unanimity the protocol effect collapses first, and the two effects converge or reverse.',
        'Define a disagreement index as |V \u2265 N| / |V \u2265 1| \u2014 the area-mean ratio of the union mask on which all raters agree \u2014 computed by the same code on all three datasets; smaller values mean greater inter-rater disagreement. Across LIDC, QUBIQ brain-growth and QUBIQ kidney the index is 0.18, 0.50 and 0.87, and the protocol-to-supervision ratio at optimal thresholds is 2.69\u00d7, 1.44\u00d7 and 0.78\u00d7 respectively. The three points are strictly monotone: the greater the rater disagreement, the more the protocol choice dominates the method choice; as agreement approaches unanimity the protocol effect collapses first, and the two effects converge or reverse. The LIDC value hides a stronger statement: the all-raters-consistent region is empty in 65% of the 792 cases, so V \u2265 4 as a reporting standard is not merely strict but undefined for most cases \u2014 an extreme protocol choice rather than a conservative one.',
        '[237] 分歧定义+65%')
    # [238] Table 7 caption
    rep(p,
        'Disagreement index: |V \u2265 N| / |V \u2265 1|, smaller meaning greater rater disagreement; ratio is protocol/supervision.',
        'Disagreement index: area-mean ratio |V \u2265 N| / |V \u2265 1|, smaller meaning greater rater disagreement; ratio is protocol/supervision.',
        '[238] caption')

doc.save(PATH)
doc2 = docx.Document(PATH)
t = doc2.paragraphs[237].text.strip()
print('\n验证 [237]:', t[:400])
