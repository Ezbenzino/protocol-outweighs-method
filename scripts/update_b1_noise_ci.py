import docx
PATH = r'outputs/paper/论文初稿_协议效应_v9.docx'
doc = docx.Document(PATH)

old = ('Pooling the four evaluation targets within each configuration, the standard deviation was '
       '0.92 and 0.53 Dice points for arm A and B on fold 0, and 0.33 and 0.27 points on fold 1, '
       'and 0.57 points pooled across both arms and both folds (2 arms \u00d7 2 folds \u00d7 4 targets '
       '\u00d7 3 seeds, 32 degrees of freedom). Taking 2\u03c3 as a practical resolution limit gives '
       '1.15 Dice points: differences smaller than this cannot be distinguished from the effect of '
       'changing the random seed alone. Two-sigma varies by evaluation target \u2014 1.29, 0.70, 1.18 '
       'and 1.31 points at V \u2265 1 to 4 under the fixed 0.5 threshold \u2014 so the pooled value of '
       '1.15 points is used as a global reference, and comparisons at a single target should be read '
       'against that target\u2019s own value.')

new = ('Pooling the four evaluation targets within each configuration, the standard deviation was '
       '0.92 and 0.53 Dice points for arm A and B on fold 0, and 0.33 and 0.27 points on fold 1, '
       'and 0.57 points pooled across both arms and both folds (2 arms \u00d7 2 folds \u00d7 4 targets '
       '\u00d7 3 seeds, 32 degrees of freedom; 95% CI [0.46, 0.76]). Taking 2\u03c3 as a practical '
       'resolution limit gives 1.15 Dice points (95% CI [0.92, 1.52]): differences smaller than this '
       'cannot be distinguished from the effect of changing the random seed alone. Two-sigma varies by '
       'evaluation target and threshold policy: under the fixed 0.5 threshold it is 1.29, 0.70, 1.18 '
       'and 1.31 points at V \u2265 1 to 4, and 0.77, 0.63, 0.56 and 0.83 points when each model '
       'selects its own optimal threshold \u2014 so the pooled fixed-threshold value of 1.15 points is '
       'used as a global reference, and comparisons at a single target should be read against that '
       'target\u2019s own value.')

found = False
for p in doc.paragraphs:
    if old in p.text:
        t = p.text
        nt = t.replace(old, new)
        if p.runs:
            p.runs[0].text = nt
            for r in p.runs[1:]:
                r.text = ''
        found = True
        print('已更新 [155]')
        break
if not found:
    print('未找到 [155] 原文!')

doc.save(PATH)

# 验证
doc2 = docx.Document(PATH)
print(doc2.paragraphs[155].text.strip()[:520])
