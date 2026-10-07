# -*- coding: utf-8 -*-
"""v21 投稿前自检（只读）。用法：python scripts/verify_paper_v21.py [docx路径]
检查：摘要词数≤250、Highlights≤85字符、正文引用按首次出现顺序编号、参考文献连续、
图表编号连续、无中文残留、列出所有待作者填写的占位符、Table 5 排序。"""
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'outputs' / 'paper' / '论文_协议效应_v21_CMIG.docx'
d = docx.Document(str(PATH))
P = [p.text for p in d.paragraphs]
ok = True


def fail(msg):
    global ok
    ok = False
    print('  ✗', msg)


# 1 摘要
i0 = P.index('Abstract')
abstract = [t for t in P[i0 + 1:i0 + 5]]
assert abstract[0].startswith('Background') and abstract[3].startswith('Conclusions')
n = sum(len(t.split()) for t in abstract)
print(f'[1] abstract words (whitespace tokens) = {n}')
if n > 250:
    fail('abstract exceeds 250 words')

# 2 Highlights
ih = P.index('Highlights')
hl = [t.lstrip('• ').strip() for t in P[ih + 1:ih + 6]]
print('[2] highlight lengths =', [len(h) for h in hl])
for h in hl:
    if len(h) > 85:
        fail(f'highlight > 85 chars: {h}')

# 3 引用顺序
iref = P.index('References')
seen, mx = [], 0
for t in P[:iref]:
    for m in re.finditer(r'\[(\d+(?:\s*[,–-]\s*\d+)*)\]', t):
        nums = []
        for part in re.split(r'\s*,\s*', m.group(1)):
            if re.search(r'[–-]', part):
                a, b = [int(x) for x in re.split(r'[–-]', part)]
                nums += list(range(a, b + 1))
            else:
                nums.append(int(part))
        for k in nums:
            if k not in seen:
                if k != mx + 1:
                    fail(f'first citation of [{k}] out of order (previous max {mx}): …{t[max(0, m.start() - 50):m.end()]}')
                seen.append(k)
                mx = max(mx, k)
refs = [t for t in P[iref + 1:] if re.match(r'^\[\d+\]', t)]
nums = [int(re.match(r'^\[(\d+)\]', t).group(1)) for t in refs]
print(f'[3] cited refs = {len(seen)} (max {mx}); reference list = {len(refs)} entries')
if nums != list(range(1, len(refs) + 1)):
    fail('reference list numbering not contiguous')
if sorted(seen) != nums:
    fail(f'cited set != listed set: missing {sorted(set(nums) - set(seen))}, extra {sorted(set(seen) - set(nums))}')

# 4 图表编号
figs = [int(m.group(1)) for t in P for m in [re.match(r'^Fig\. (\d+) —', t)] if m]
tabs = [int(m.group(1)) for t in P for m in [re.match(r'^Table (\d+)\.', t)] if m]
print('[4] figure captions', figs, '| table captions', tabs, '| tables in body', len(d.tables))
if figs != list(range(1, 10)) and sorted(figs) != list(range(1, 10)):
    fail('figure captions not 1–9')
if sorted(tabs) != list(range(1, 10)):
    fail('table captions not 1–9')

# 4b 图表按首次引用顺序编号（排除图注/表题本身与 "their Table/Fig"）
firsts = {'Fig': [], 'Table': []}
for t in P[:iref]:
    if re.match(r'^(Fig\. \d+ —|Table \d+\.)', t):
        continue
    for m in re.finditer(r'(Fig|Table)s?\.? (\d+)', t):
        ctx = t[max(0, m.start() - 45):m.start()]
        if re.search(r"[Tt]heir|Zhi et al\.’?'?s", ctx):
            continue  # 指 Zhi et al. 的图表，不是本文的
        kind, k = m.group(1), int(m.group(2))
        if k not in firsts[kind]:
            firsts[kind].append(k)
print('[4b] first-citation order  Fig:', firsts['Fig'], '| Table:', firsts['Table'])
for kind in firsts:
    if firsts[kind] != sorted(firsts[kind]):
        fail(f'{kind} numbers are not in order of first citation')

# 5 中文残留 / 占位符
cjk = [(i, t[:60]) for i, t in enumerate(P) if re.search(r'[一-鿿]', t)]
for c in cjk:
    fail(f'CJK text remains in paragraph {c[0]}: {c[1]}')
ph = [(i, m.group(0)) for i, t in enumerate(P) for m in re.finditer(r'\[(?:to be completed|PENDING[^\]]*|Funding statement to be completed[^\]]*|Code repository URL[^\]]*)\]', t)]
print('[5] placeholders left for the authors:')
for i, s in ph:
    print('     P%d: %s' % (i, s[:90]))

# 6 Table 5 排序
t5 = d.tables[4]
vals = []
for r in t5.rows[1:]:
    m = re.match(r'([\d.]+)', r.cells[2].text.strip())
    vals.append((float(m.group(1)), r.cells[0].text[:45], r.cells[1].text))
effects = [v for v in vals if v[2] != 'Noise floor']
print('[6] Table 5 order:', [v[0] for v in vals])
if [v[0] for v in effects] != sorted([v[0] for v in effects], reverse=True):
    fail('Table 5 effects not in descending order')

# 7 已知易错表述
bad = ['Spearman rho', 'rho =', 'five prior methods', 'survey the field', 'hierarchical extension',
       'two lesion radii', 'encoder learning rate', 'costs 27 Dice', 'data-preparation and learning-design effects',
       'exceeds every data-preparation', 'over [0, 1]', '4696–4705', 'between 82.3 and 86.4']
for b in bad:
    for i, t in enumerate(P):
        if b in t:
            fail(f'stale wording "{b}" in P{i}')
print('\nRESULT:', 'ALL CHECKS PASSED' if ok else 'SOME CHECKS FAILED')
