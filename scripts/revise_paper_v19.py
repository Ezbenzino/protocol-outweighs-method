# -*- coding: utf-8 -*-
"""revise_paper_v19.py —— v18 -> v19：把 §5.5 里"没有 nnU-Net"这半句补成完整论证。

问题：§5.5 该段标题是 "Two dimensions only, and no nnU-Net"，但正文只讲了 2D 对 3D，
**完全没有解释为什么没有 nnU-Net**。标题许诺了一个解释，身子里没有——
审稿人扫到这里会认为作者知道该做却回避了。

修法：把它写成一个主动的设计决策，并引用架构轴上现有的最大规模证据
（Zhi et al. [31] 固定协议下重训 23 个模型，五个最好的差 0.31 点；
nnU-Net Revisited [28] 报告许多架构优势在协议对齐后不成立），
两者方向都与本文的架构对照一致。这比自己再跑一个 nnU-Net 更有说服力。
"""
import os
import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = os.path.join(ROOT, 'outputs', 'paper')
SRC = os.path.join(P, '论文_协议效应_v18_CIBM.docx')
DST = os.path.join(P, '论文_协议效应_v19_CIBM.docx')

OLD = ('Two dimensions only, and no nnU-Net. All arms are 2-D, for compute reasons. 2-D '
       'evaluation is itself a protocol choice, and this is a tension we acknowledge rather than '
       'resolve. Zhi et al. report that 2-D models score systematically higher than 3-D ones '
       '[31], which may say more about the relative difficulty of the two evaluation protocols '
       'than about the models.')

NEW = ('Two dimensions only, and no self-configuring framework. All arms are 2-D, for compute '
       'reasons. 2-D evaluation is itself a protocol choice, and this is a tension we acknowledge '
       'rather than resolve. Zhi et al. report that 2-D models score systematically higher than '
       '3-D ones [31], which may say more about the relative difficulty of the two evaluation '
       'protocols than about the models. We also did not run a self-configuring pipeline such as '
       'nnU-Net, and this follows from the design rather than being an oversight: such a pipeline '
       'varies architecture, patch size, augmentation, loss, post-processing and ensembling '
       'together, so its result cannot enter a table whose entries are defined by changing one '
       'factor at a time, and would constitute a robustness check rather than a controlled '
       'contrast. On the architecture axis itself the largest available evidence is not ours. Zhi '
       'et al. retrained 23 published models under one fixed protocol and found 0.31 points '
       'between their five best across a 558-fold parameter range [31], and nnU-Net Revisited '
       'reports that many claimed architectural advances do not survive protocol-matched '
       'retraining [28]. Both point the same way as the architecture control reported here.')


def main():
    doc = docx.Document(SRC)
    hit = False
    for par in doc.paragraphs:
        joined = ''.join(r.text for r in par.runs)
        if OLD in joined:
            par.runs[0].text = joined.replace(OLD, NEW)
            for r in par.runs[1:]:
                r.text = ''
            hit = True
            break
    assert hit, '未找到 §5.5 的 2D/nnU-Net 段'
    doc.save(DST)

    d2 = docx.Document(DST)
    txt = '\n'.join(p.text for p in d2.paragraphs)
    assert 'Two dimensions only, and no nnU-Net.' not in txt, '旧标题仍在'
    for good in ('no self-configuring framework', 'follows from the design rather than being an '
                 'oversight', 'retrained 23 published models under one fixed protocol and found '
                 '0.31 points', 'nnU-Net Revisited reports that many claimed architectural '
                 'advances do not survive'):
        assert good in txt, f'缺少: {good}'
    a = [i for i, p in enumerate(d2.paragraphs) if p.text.strip() == 'Abstract'][0]
    h = [i for i, p in enumerate(d2.paragraphs) if p.text.strip() == 'Highlights'][0]
    n = sum(len(p.text.split()) for p in d2.paragraphs[a + 1:h]
            if p.text.strip() and not p.text.strip().startswith('Keywords'))
    print(f'摘要 {n} 词；段落 {len(d2.paragraphs)}  表 {len(d2.tables)}  图 {len(d2.inline_shapes)}  '
          f"PENDING {txt.count('PENDING')}")
    print(f'正文词数 {sum(len(p.text.split()) for p in d2.paragraphs)}')
    print('回读自检通过 ->', DST)


if __name__ == '__main__':
    main()
