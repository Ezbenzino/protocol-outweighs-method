from pptx import Presentation
from pptx.util import Inches, Pt
import json
import sys

if len(sys.argv) < 2:
    raise SystemExit("用法: python scripts/_read_ppt.py <pptx路径>")
prs = Presentation(sys.argv[1])
print(f'总页数: {len(prs.slides)}')
print(f'幻灯片尺寸: {prs.slide_width/914400:.1f} x {prs.slide_height/914400:.1f} 英寸')
print('='*80)

for i, slide in enumerate(prs.slides):
    print(f'\n===== 第 {i+1} 页 =====')
    print(f'布局: {slide.slide_layout.name}')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    sizes = []
                    for run in para.runs:
                        if run.font.size:
                            sizes.append(run.font.size.pt)
                    size_str = f' [字号:{sizes[0]}]' if sizes else ''
                    print(f'  {text}{size_str}')
        if shape.has_table:
            print('  [表格]')
            table = shape.table
            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip().replace(chr(10),' ') for cell in row.cells]
                print(f'    行{row_idx}: ' + ' | '.join(cells))
