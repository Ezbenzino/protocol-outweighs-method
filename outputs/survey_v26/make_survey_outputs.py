# -*- coding: utf-8 -*-
"""Regenerate survey CSVs, the supplementary spreadsheet and stats.json from epmc_raw.json, decisions.json, codes.json."""
import collections, csv, html, json, re
from openpyxl import Workbook
from openpyxl.styles import Font
raw = json.load(open('epmc_raw.json')); R = raw['results']
dec = json.load(open('decisions.json')); codes = json.load(open('codes.json'))

def meta(x):
    j = x.get('journalInfo', {})
    jt = (j.get('journal', {}) or {}).get('medlineAbbreviation') or (j.get('journal', {}) or {}).get('title') or ''
    fa = (x.get('authorString', '') or '').split(',')[0].strip().rstrip('.')
    return dict(first_author=fa, year=x.get('pubYear'), journal=jt, doi=x.get('doi', ''), pmid=x.get('pmid', ''),
                pmcid=x.get('pmcid', ''), title=html.unescape(re.sub('<[^>]+>', '', x.get('title', ''))).strip())

with open('S2a_survey_records.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['record', 'source', 'first_author', 'year', 'journal', 'title', 'doi', 'pmid', 'pmcid', 'decision', 'reason'])
    for i, x in enumerate(R):
        m = meta(x); d, r = dec[str(i)]
        w.writerow([f'{i:03d}', x.get('source'), m['first_author'], m['year'], m['journal'], m['title'], m['doi'], m['pmid'], m['pmcid'], d, r])
CROP = {'C': 'located by reference annotation', 'U': 'not located by reference annotation', 'NS': 'not stated'}
EVC = {'Y': 'yes', 'N': 'no', 'NS': 'not stated'}; YN = {'S': 'stated', 'N': 'not stated'}
with open('S2b_survey_coding.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['record', 'first_author', 'year', 'journal', 'doi', 'pmcid', 'datasets', 'dimensionality', 'headline_overlap',
                'fusion_rule', 'fusion_detail', 'threshold', 'threshold_detail', 'training_crop', 'training_crop_detail',
                'evaluation_extent', 'evaluation_extent_detail', 'evaluated_in_annotation_located_crop', 'all_four_reported',
                'evidence_fusion', 'evidence_threshold', 'evidence_crop', 'evidence_extent', 'note'])
    for k in sorted(codes):
        e = codes[k]; m = meta(R[int(k)])
        all4 = 'yes' if (e['fus'] == 'S' and e['thr'] == 'S' and e['crop'] in ('C', 'U') and e['ext'] == 'S') else 'no'
        w.writerow([k, m['first_author'], m['year'], m['journal'], m['doi'], m['pmcid'], e.get('data', ''), e.get('dim', ''), e.get('dice', ''),
                    YN[e['fus']], e.get('fus_rule', ''), YN[e['thr']], e.get('thr_val', ''), CROP[e['crop']], e.get('crop_detail', ''),
                    YN[e['ext']], e.get('ext_detail', ''), EVC[e['evc']], all4,
                    ' | '.join(e.get('fus_ev', [])), ' | '.join(e.get('thr_ev', [])), ' | '.join(e.get('crop_ev', [])),
                    ' | '.join(e.get('ext_ev', [])), (e.get('note', '') + (' ; ' + e['evc_detail'] if e.get('evc_detail') else '')).strip(' ;')])

C = collections.Counter(v[0] for v in dec.values()); n = len(codes); pct = lambda k: round(100 * k / n)
c = lambda f, v: sum(1 for e in codes.values() if e[f] == v)
S = dict(query='the query TITLE_ABS:("LIDC-IDRI" OR "LIDC" OR "LUNA16" OR "LUNA") AND TITLE_ABS:(segment*) AND '
               'TITLE_ABS:(nodule*) AND PUB_YEAR:[2018 TO 2026]', date='15 October 2026', n_records=len(dec),
         n_ppr=C['excluded: preprint'], n_dup=C['excluded: duplicate record'], n_abs_excl=C['excluded at title/abstract'],
         n_ft_excl=C['excluded at full text'], n_coded=C['eligible; full text read and coded'],
         n_noft=C['eligible; full text not available in PMC'], n_ft_blocked=C['eligible; PMC full text not retrievable at survey time'])
S['n_screened'] = S['n_records'] - S['n_ppr'] - S['n_dup']; S['n_candidates'] = S['n_screened'] - S['n_abs_excl']
S['n_ft_read'] = S['n_coded'] + S['n_ft_excl']; S['n_eligible'] = S['n_coded'] + S['n_noft'] + S['n_ft_blocked']
for f, v, k in [('thr', 'N', 'thr_N'), ('fus', 'N', 'fus_N'), ('ext', 'N', 'ext_N'), ('crop', 'C', 'crop_A'), ('crop', 'U', 'crop_O'),
                ('crop', 'NS', 'crop_NS'), ('evc', 'Y', 'evc_Y')]:
    S[k] = c(f, v); S[k + '_pct'] = pct(S[k])
S['all4'] = sum(1 for e in codes.values() if e['fus'] == 'S' and e['thr'] == 'S' and e['crop'] in ('C', 'U') and e['ext'] == 'S')
S['all4_pct'] = pct(S['all4'])
S['none4'] = sum(1 for e in codes.values() if e['fus'] == 'N' and e['thr'] == 'N' and e['crop'] == 'NS' and e['ext'] == 'N')
S['fus_N_strict'] = S['fus_N'] + sum(1 for k in ('042', '099', '132') if codes[k]['fus'] == 'S'); S['fus_N_strict_pct'] = pct(S['fus_N_strict'])
json.dump(S, open('stats.json', 'w'), indent=1)

wb = Workbook(); ws = wb.active; ws.title = 'README'
lines = [
    'Supplementary data — literature survey of reporting practice in LIDC-IDRI / LUNA16 lung nodule segmentation',
    'Manuscript: Protocol outweighs method (Physics in Medicine & Biology submission, v26)', '',
    'Database: Europe PMC REST API (https://www.ebi.ac.uk/europepmc/webservices/rest/search), resultType=core',
    'Query: ' + raw['query'], 'Search run (UTC): ' + raw['retrieved'][:19].replace('T', ' '), f"Records returned: {raw['hitCount']}",
    f"Flow: {S['n_records']} records -> {S['n_ppr']} preprints and {S['n_dup']} duplicate records removed -> {S['n_screened']} screened on title/abstract -> {S['n_abs_excl']} excluded -> {S['n_candidates']} candidates",
    f"      {S['n_candidates']} candidates -> {S['n_ft_read']} with retrievable PubMed Central full text ({S['n_ft_excl']} excluded on reading, {S['n_coded']} coded); {S['n_ft_blocked']} with PMC full text not retrievable; {S['n_noft']} without PMC full text",
    f"      eligible studies = {S['n_eligible']} ({S['n_coded']} coded + {S['n_eligible'] - S['n_coded']} retained but never counted in a denominator)",
    'Eligibility: trains or applies a learning-based model and reports a Dice or IoU score for lung nodule segmentation on LIDC-IDRI or LUNA16;',
    '             detection-only, classification-only, synthesis, dataset, lung-parenchyma, review and non-learning studies excluded; preprints excluded.',
    'Full text: Europe PMC fullTextXML or NCBI PMC rendering. Quotations are taken from those renderings; inline mathematics may appear as plain text.',
    'Coded fields (sheet "coding"):',
    "  fusion_rule — is the rule that turns the readers' contours (or other source) into the reference mask stated?",
    '  threshold — is the binarisation threshold (value, argmax, level set or selection procedure) stated for the final mask?',
    '  training_crop — are training inputs crops located by the reference annotation, not so located (whole slice/volume, random, detector-supplied), or not stated?',
    '  evaluation_extent — can the image region over which the overlap score is computed be determined from the text?',
    '  evaluated_in_annotation_located_crop — is the test-time evaluation region located by the reference annotation (localisation assumed solved)?',
    '  all_four_reported — fusion rule, threshold, training-crop rule and evaluation extent all stated.',
    'Rater: single rater. Each coded field is accompanied by the verbatim supporting passage(s).',
]
for i, l in enumerate(lines, 1): ws.cell(i, 1, l)
ws['A1'].font = Font(bold=True); ws.column_dimensions['A'].width = 160
for name, fn in [('records', 'S2a_survey_records.csv'), ('coding', 'S2b_survey_coding.csv')]:
    s = wb.create_sheet(name)
    for r in csv.reader(open(fn, encoding='utf-8-sig')): s.append(r)
    for cc in s[1]: cc.font = Font(bold=True)
    s.freeze_panes = 'B2'
    for col in s.columns:
        L = max(len(str(x.value or '')) for x in col[:60]); s.column_dimensions[col[0].column_letter].width = min(60, max(8, L * 0.9))
wb.save('Supplementary_data_S2_literature_survey.xlsx')
print({k: S[k] for k in ('thr_N', 'thr_N_pct', 'fus_N_pct', 'ext_N_pct', 'crop_A_pct', 'evc_Y_pct', 'all4', 'all4_pct', 'none4', 'fus_N_strict_pct')})
